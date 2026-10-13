"""Opt-in, single-batch legacy deployment conversion. Never replaces masters/history."""
from datetime import date, datetime, timedelta
from io import BytesIO
import json
from uuid import UUID, uuid4

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError, OperationalError

from app.core.config import settings
from app.domain.capabilities import require_capabilities
from app.domain.holidays import business_days_between, month_workday_summary
from app.domain.kpi import pmo_monthly_kpi_values
from app.domain.people import is_pmo_personnel
from app.enums import Capability, ProjectLogStatus
from app.models.core import Personnel, Project, ProjectAssignment, ProjectLog, MonthlyEmploymentMM, MonthlyKpiSummary
from app.schemas.assignments import AssignmentCreate
from app.services import data_backup as backup
from app.services.excel_common import MAX_FILE_BYTES as MAX_FILE_BYTES, MAX_ROWS, read_rows
from app.services.personnel_master import refresh_personnel_projections, refresh_open_kpi
from app.services.project_assignments import EDIT_FIELDS, VERSION_MODELS, _values, refresh_assignment_mm, version

SHEET = "최초투입전환"
COLUMNS = [("배정키", "source_key"), ("프로젝트 코드", "project_code"), ("personnel_id", "personnel_id"),
           ("배정 유형", "assignment_type"), ("프로젝트 역할", "assignment_role"), ("배정 상태", "assignment_status"),
           ("시작일", "start_date"), ("종료일", "end_date"), ("월 기준 MM", "monthly_base_mm"),
           ("확도율", "certainty_rate"), ("상주 유형", "onsite_type"), ("대표 여부", "is_primary"),
           ("표시 순서", "sequence_no"), ("비고", "note")]
LABELS = dict((field, label) for label, field in COLUMNS)
MAPS = {
    "assignment_type": {"수행": "delivery", "제안": "proposal", "지원": "support", "미투입": "unassigned"},
    "assignment_status": {"예정": "planned", "투입": "assigned", "종료": "ended", "취소": "cancelled"},
    "assignment_role": {"제안PM": "proposal_pm", "발표PM": "presentation_pm", "수행PM": "delivery_pm", "제안팀": "proposal_team", "수행팀": "delivery_team", "지원팀": "support_team"},
    "onsite_type": {"상주": "onsite", "비상주": "remote", "혼합": "hybrid"},
}
MARKER = "[assignment_import_complete]"
VALIDATIONS: dict[str, dict] = {}
TTL = timedelta(minutes=30)
INPUT_RULES = {
    "monthly_base_mm": "0~9999.99의 숫자를 소수 둘째 자리까지 입력해 주세요.",
    "certainty_rate": "0~100의 숫자를 소수 둘째 자리까지 입력해 주세요.",
    "sequence_no": "0~100000의 정수를 입력해 주세요.",
    "start_date": "필수 날짜를 YYYY-MM-DD로 입력해 주세요.",
    "end_date": "필수 날짜를 YYYY-MM-DD로 입력해 주세요.",
    "note": "비고는 4000자 이내의 텍스트로 입력해 주세요.",
}


def completed(session):
    # Global system log has no project and cannot be edited via project-log CRUD.
    return session.scalar(select(ProjectLog).where(ProjectLog.project_id.is_(None), ProjectLog.content.startswith(MARKER, autoescape=True)).limit(1))


def status(session, user):
    require_capabilities(session, user, Capability.ASSIGNMENT_MANAGE)
    row = completed(session)
    enabled = settings.assignment_import_enabled and settings.app_environment != "production"
    return {"enabled": enabled, "completed": row is not None,
            "message": "이 데이터의 최초 전환은 이미 완료되었습니다." if row else "최초 전환 도구가 활성화되어 있습니다." if enabled else "최초 전환 도구는 비활성화 상태입니다.",
            "completion": json.loads(row.content[len(MARKER):]) if row else None}


def require_open(session, user):
    current = status(session, user)
    if not current["enabled"]:
        raise HTTPException(status_code=403, detail="최초 전환 도구를 별도 전환 환경에서 활성화해 주세요. 운영 환경에서는 사용할 수 없습니다.")
    if current["completed"]:
        raise HTTPException(status_code=409, detail=current["message"])


def _workbook_bytes(workbook):
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EAF0FA")
            sheet.column_dimensions[cell.column_letter].width = 38 if cell.value == "personnel_id" else 24
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Preserve literal notes/names, never Excel formulas.
    raw = BytesIO()
    workbook.save(raw)
    return raw.getvalue()


def template(session, user, *, references=False):
    require_open(session, user)
    workbook = Workbook()
    sheet = workbook.active
    if references:
        sheet.title = "프로젝트 참조"
        sheet.append(["프로젝트 코드", "사업명", "사용 여부"])
        for p in session.scalars(select(Project).order_by(Project.code)).all():
            sheet.append([p.code, p.name, bool(p.project_code and p.project_code.is_active)])
        sheet = workbook.create_sheet("인력 참조")
        sheet.append(["personnel_id", "이름", "사번", "인력 유형", "본부", "팀", "재직 상태", "사용 여부"])
        for p in session.scalars(select(Personnel).order_by(Personnel.name, Personnel.id)).all():
            sheet.append([p.id, p.name, p.employee_no, p.personnel_type.value, p.group_name, p.team_name, p.employment_status.value, p.is_active])
        sheet = workbook.create_sheet("입력 안내")
        sheet.append(["항목", "입력 기준"])
        sheet.append(["배정키", "원본 행을 구분하는 고유 키. 같은 파일 내 중복 불가"])
        sheet.append(["월 기준 MM", "필수 0~9999.99, 소수 둘째 자리. 총/누적/당월 MM을 복사하지 말고 의미 확인"])
        sheet.append(["기간", "시작일·종료일 필수 YYYY-MM-DD. 종료 상태는 실제 종료일까지"])
        sheet.append(["확도율", "0~100, 빈 값은 100%. 소수 둘째 자리"])
        sheet.append(["대표 여부 / 표시 순서", "참/거짓 또는 1/0 (기본 거짓), 정수 0~100000 (기본 0)"])
        for field, values in MAPS.items():
            sheet.append([LABELS[field], " / ".join(f"{label}({code})" for label, code in values.items())])
        sheet.append(["반영", "한 파일 한 번만 추가. 기존 데이터 수정/삭제 없음. 과거 스냅샷/KPI 보존"])
    else:
        sheet.title = SHEET
        sheet.append([label for label, _ in COLUMNS])
    return _workbook_bytes(workbook)


def _parse(cells, projects):
    values = {field: cell.value for (_, field), cell in zip(COLUMNS, cells)}
    for cell in cells:
        if cell.data_type == "f":
            raise ValueError(f"수식은 허용하지 않습니다: {cell.coordinate}")
    for key, value in values.items():
        if isinstance(value, str):
            values[key] = value.strip() or None
    source = values.pop("source_key")
    code = values.pop("project_code")
    if not isinstance(source, str) or len(source) > 200:
        raise ValueError("배정키는 200자 이내의 필수 텍스트입니다.")
    if code not in projects:
        raise ValueError("프로젝트 코드가 없습니다. 프로젝트 마스터를 먼저 등록해 주세요.")
    if projects[code].project_code is None or projects[code].project_code.code != code:
        raise ValueError("프로젝트와 프로젝트 코드 연결을 먼저 보정해 주세요.")
    try:
        UUID(str(values["personnel_id"]))
    except (ValueError, TypeError) as exc:
        raise ValueError("personnel_id는 인력 다운로드에 있는 UUID를 사용해 주세요.") from exc
    values["project_id"] = projects[code].id
    for field, labels in MAPS.items():
        values[field] = labels.get(values[field], values[field])
    for field in ("start_date", "end_date"):
        if isinstance(values[field], datetime):
            values[field] = values[field].date()
    primary = values["is_primary"]
    if primary not in (None, True, False, 0, 1, "참", "거짓", "1", "0", "true", "false"):
        raise ValueError("대표 여부는 참/거짓 또는 1/0으로 입력해 주세요.")
    values["is_primary"] = primary in (True, 1, "참", "1", "true")
    values["sequence_no"] = values["sequence_no"] if values["sequence_no"] is not None else 0
    return source, code, AssignmentCreate.model_validate(values)


def validate_workbook(session, user, file_name, raw):
    require_open(session, user)
    rows = read_rows(file_name, raw, SHEET)
    if not rows or [cell.value for cell in rows[0]] != [label for label, _ in COLUMNS]:
        raise HTTPException(status_code=400, detail="최초투입전환 템플릿의 컬럼과 순서를 그대로 사용해 주세요.")
    if len(rows) - 1 > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"최대 {MAX_ROWS}행까지 가능합니다.")
    original_version = version(session)
    project_rows = session.scalars(select(Project)).all()
    projects = {p.code: p for p in project_rows}
    if len(projects) != len(project_rows):
        raise HTTPException(status_code=409, detail="프로젝트 코드 중복을 먼저 해결해 주세요.")
    issues, previews, payloads, warnings = [], [], [], []
    keys, signatures, by_person = set(), set(), {}
    for number, cells in enumerate(rows[1:], 2):
        if all(cell.value is None for cell in cells):
            continue
        try:
            source, code, payload = _parse(cells, projects)
            _, values, row_warnings, _, _ = _values(session, payload, historical_import=True)
            person = session.get(Personnel, payload.personnel_id)
            if is_pmo_personnel(person) and not person.mm_start_date and values["assignment_status"] != "cancelled" and values["assignment_type"] != "unassigned":
                raise ValueError("PMO 인력의 MM 시작일을 먼저 보완해 주세요. 과거 재직 MM 분모를 추정하지 않습니다.")
            if is_pmo_personnel(person) and ((person.mm_start_date and values["start_date"] < person.mm_start_date) or (person.mm_end_date and values["end_date"] > person.mm_end_date)):
                row_warnings.append("배정 기간이 인력의 MM 재직 기간 밖에 걸칩니다. 원본 기간과 MM 시작·종료일을 대사해 주세요. 기간이나 산식을 자동 보정하지 않습니다.")
            signature = tuple(values.get(key) for key in ("project_id", "personnel_id", "assignment_type", "assignment_role", "start_date", "end_date"))
            if source in keys or signature in signatures:
                raise ValueError("파일 안에 배정키 또는 완전 중복 배정이 있습니다.")
            # Also block exact historical/cancelled duplicates omitted by normal operational overlap checks.
            existing = session.scalars(select(ProjectAssignment).where(ProjectAssignment.personnel_id == values["personnel_id"], ProjectAssignment.project_id == values["project_id"])).all()
            if any(tuple(getattr(row, key) for key in ("project_id", "personnel_id", "assignment_type", "assignment_role", "start_date", "end_date")) == signature for row in existing):
                raise ValueError("기존 데이터와 완전 중복인 배정입니다. 기존 배정은 덮어쓰지 않습니다.")
            operational = values["assignment_status"] != "cancelled" and values["assignment_type"] != "unassigned"
            previous = by_person.setdefault(values["personnel_id"], [])
            if operational and any(old["assignment_status"] != "cancelled" and old["assignment_type"] != "unassigned" and old["start_date"] <= values["end_date"] and old["end_date"] >= values["start_date"] for old in previous):
                row_warnings.append("같은 파일의 다른 배정과 기간이 겹칩니다. 원본의 중복 투입을 확인해 주세요.")
            keys.add(source)
            signatures.add(signature)
            previous.append(values)
            warnings.extend(f"{number}행: {message}" for message in row_warnings)
            previews.append({"row_number": number, "source_key": source, "project_code": code, "personnel_name": person.name, "personnel_type": person.personnel_type.value, "assignment_status": values["assignment_status"], "monthly_base_mm": float(payload.monthly_base_mm), "warnings": row_warnings})
            payloads.append({"source_key": source, "payload": payload})
        except (ValueError, HTTPException, ValidationError) as exc:
            if isinstance(exc, ValidationError):
                message = " / ".join(f"{LABELS.get(str(e['loc'][0]), str(e['loc'][0]))}: {INPUT_RULES.get(str(e['loc'][0]), '입력 안내에 있는 유효한 코드/라벨을 입력해 주세요.')}" for e in exc.errors())
            else:
                message = str(exc.detail if isinstance(exc, HTTPException) else exc)
            issues.append({"row_number": number, "message": message})
    if not previews and not issues:
        raise HTTPException(status_code=400, detail="반영할 배정이 없습니다.")
    if version(session) != original_version:
        raise HTTPException(status_code=409, detail="검증 중 데이터가 변경되었습니다. 다시 검증해 주세요.")
    now = backup._now()
    for key, cached in list(VALIDATIONS.items()):
        if cached["created_at"] + TTL < now:
            VALIDATIONS.pop(key, None)
    if len(VALIDATIONS) >= 100:
        raise HTTPException(status_code=429, detail="검증 요청이 많습니다. 잠시 후 다시 시도해 주세요.")
    token = str(uuid4())
    result = {"validation_id": token, "file_sha256": backup._file_sha256(raw), "file_name": file_name, "summary": {"new": len(previews), "error": len(issues)}, "rows": previews, "issues": issues, "warnings": warnings}
    VALIDATIONS[token] = {"created_at": now, "user_id": user.id, "version": original_version, "result": result, "payloads": payloads}
    return result


def _fill_missing_projections(session, imported, horizon):
    today = date.today()
    start = min((row.start_date.year, row.start_date.month) for row in imported)
    people = session.scalars(select(Personnel)).all()
    pmo_people = [person for person in people if is_pmo_personnel(person) and person.mm_start_date]
    start = min([start, *((person.mm_start_date.year, person.mm_start_date.month) for person in pmo_people)])
    horizon = max([horizon, *((person.mm_end_date.year, person.mm_end_date.month) for person in pmo_people if person.mm_end_date)])
    if horizon[0] > today.year + 100 or horizon[0] - start[0] > 200:
        raise HTTPException(status_code=409, detail="재직 MM 계산 기간이 너무 깁니다. 인력 MM 시작·종료일을 먼저 확인해 주세요.")
    employment = {}
    for row in session.scalars(select(MonthlyEmploymentMM)).all():
        key = (row.personnel_id, row.year, row.month)
        if key in employment:
            raise HTTPException(status_code=409, detail="같은 인력/월의 재직 MM 원장이 중복되어 먼저 확인이 필요합니다.")
        employment[key] = row
    kpis = {}
    for row in session.scalars(select(MonthlyKpiSummary).where(MonthlyKpiSummary.organization_name == "PMO본부")).all():
        key = (row.year, row.month)
        if key in kpis:
            raise HTTPException(status_code=409, detail="같은 조직/월의 KPI 원장이 중복되어 먼저 확인이 필요합니다.")
        kpis[key] = row
    counts = {"employment_created": 0, "past_kpi_created": 0, "past_kpi_preserved": 0}
    year, month = start
    while (year, month) <= horizon:
        workdays = int(month_workday_summary(session, year=year, month=month)["workdays"])
        for person in people:
            if not is_pmo_personnel(person) or not person.mm_start_date or (person.id, year, month) in employment:
                continue
            if not person.is_active and (year, month) >= (today.year, today.month):
                continue
            window = backup._person_employment_window(person, year, month, today)
            if window:
                days = business_days_between(session, *window)
                session.add(MonthlyEmploymentMM(personnel_id=person.id, year=year, month=month, workdays=workdays, employed_workdays=days, employment_mm=round(days / workdays if workdays else 0, 4)))
                counts["employment_created"] += 1
        session.flush()
        if (year, month) < (today.year, today.month):
            if (year, month) in kpis:
                counts["past_kpi_preserved"] += 1
            else:
                session.add(MonthlyKpiSummary(year=year, month=month, organization_name="PMO본부", source_snapshot_date=today, **pmo_monthly_kpi_values(session, year, month)))
                counts["past_kpi_created"] += 1
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    for person in people:
        refresh_personnel_projections(session, person, employment=False)
    refresh_open_kpi(session, through=horizon)
    return counts


def apply_workbook(session, user, validation_id, file_sha256, acknowledge_warnings):
    require_open(session, user)
    cached = VALIDATIONS.get(validation_id)
    if not cached or cached["created_at"] + TTL < backup._now():
        raise HTTPException(status_code=409, detail="검증 결과가 없거나 만료되었습니다. 다시 검증해 주세요.")
    if cached["user_id"] != user.id:
        raise HTTPException(status_code=403, detail="다른 사용자의 검증은 반영할 수 없습니다.")
    result = cached["result"]
    if file_sha256 != result["file_sha256"] or result["issues"]:
        raise HTTPException(status_code=409, detail="검증 파일이 다르거나 오류가 있습니다.")
    if not acknowledge_warnings:
        raise HTTPException(status_code=409, detail="1회 반영과 과거 데이터 보존·주의사항을 확인해 주세요.")
    backup._acquire_lock(session, locked_by=user.name)
    before = None
    try:
        if session.get_bind().dialect.name == "sqlite":
            session.connection().exec_driver_sql("BEGIN IMMEDIATE")
        else:
            for model in VERSION_MODELS:
                session.scalars(select(model).order_by(model.id).with_for_update()).all()
        require_open(session, user)
        if VALIDATIONS.get(validation_id) is not cached or version(session) != cached["version"]:
            raise HTTPException(status_code=409, detail="검증 후 DB가 변경되었습니다. 다시 검증해 주세요.")
        before = backup.create_backup(session, actor_name=user.name, kind="assignment_import_before_backup", source_file_name=result["file_name"])
        imported, ids = [], []
        horizon = (date.today().year, 12)
        for item in cached["payloads"]:
            _, values, _, _, _ = _values(session, item["payload"], historical_import=True)
            row = ProjectAssignment(**values)
            session.add(row)
            session.flush()
            horizon = max(horizon, refresh_assignment_mm(session, row, historical_import=True))
            imported.append(row)
            ids.append({"source_key": item["source_key"], "assignment_id": row.id})
            session.add(ProjectLog(project_id=row.project_id, log_status=ProjectLogStatus.MEMO, logged_at=backup._now(), author_name=user.name, updated_by_name=user.name,
                                   content="최초 투입 전환 " + json.dumps({"source_key": item["source_key"], "assignment_id": row.id, "file_sha256": file_sha256, "values": {key: backup._serialize_scalar(values[key]) for key in EDIT_FIELDS}}, ensure_ascii=False)))
            row.project.recent_activity_at = max(row.project.recent_activity_at or backup._now(), backup._now())
        counts = {"assignments_created": len(imported), **_fill_missing_projections(session, imported, horizon)}
        completion = {"file_sha256": file_sha256, "file_name": result["file_name"], "completed_at": backup._now().isoformat(), "actor_name": user.name, "pre_backup_id": before["backup_id"], "counts": counts}
        session.add(ProjectLog(project_id=None, log_status=ProjectLogStatus.MEMO, logged_at=backup._now(), author_name=user.name, updated_by_name=user.name, content=MARKER + json.dumps(completion, ensure_ascii=False)))
        session.commit()
    except Exception as exc:
        session.rollback()
        try:
            backup._append_manifest({"id": backup._make_id("OPR"), "created_at": backup._now().isoformat(), "kind": "assignment_import_failed", "status": "failed", "actor_name": user.name, "source_backup_id": before["backup_id"] if before else None, "message": str(exc)})
        except OSError:
            pass
        if isinstance(exc, (IntegrityError, OperationalError)):
            raise HTTPException(status_code=409, detail="반영 중 DB 충돌이 발생해 전체 롤백했습니다. 다시 검증해 주세요.") from exc
        raise
    finally:
        backup._release_lock(session)
    VALIDATIONS.pop(validation_id, None)
    message = "최초 투입 전환을 완료했습니다. 추가 전환은 차단됩니다."
    try:
        backup._append_manifest({"id": backup._make_id("OPR"), "created_at": backup._now().isoformat(), "kind": "assignment_import_complete", "status": "success", "source_backup_id": before["backup_id"], "actor_name": user.name, "counts": counts})
    except OSError:
        message += " 작업 이력 파일 기록은 실패했으나 DB 완료 기록과 사전 백업은 보존됩니다."
    return {**completion, "created_ids": ids, "message": message}
