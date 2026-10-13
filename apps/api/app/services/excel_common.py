"""Shared workbook parsing and deterministic DB change detection."""
from datetime import datetime
from io import BytesIO
import json

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.services import data_backup as backup

MAX_ROWS = 5000
MAX_FILE_BYTES = 20 * 1024 * 1024


def scalar(value):
    return value.isoformat() if isinstance(value, datetime) else backup._serialize_scalar(value)


def fingerprint(session, models):
    payload = {
        model.__tablename__: [
            {column.name: scalar(getattr(row, column.name)) for column in model.__table__.columns}
            for row in session.scalars(select(model).order_by(*model.__table__.primary_key.columns).execution_options(populate_existing=True)).all()
        ] for model in models
    }
    return backup._file_sha256(json.dumps(payload, sort_keys=True, ensure_ascii=True).encode())


def read_rows(file_name, raw, sheet_name):
    if not file_name.lower().endswith(".xlsx") or not raw or len(raw) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="20MB 이하의 .xlsx 파일을 선택해 주세요.")
    workbook = None
    try:
        workbook = load_workbook(BytesIO(raw), data_only=False, read_only=True)
        if workbook.sheetnames != [sheet_name]:
            raise ValueError(f"{sheet_name} 단일 시트 파일이 필요합니다.")
        sheet = workbook[sheet_name]
        if sheet.max_row > MAX_ROWS + 4 or sheet.max_column > 60:
            raise ValueError(f"파일은 {MAX_ROWS}행 이내로 입력해 주세요.")
        return list(sheet.iter_rows())
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"엑셀을 읽을 수 없습니다: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
