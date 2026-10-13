"""Personnel-only workbook import using persisted UUIDs and explicit types."""
from datetime import datetime, timedelta
from io import BytesIO
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError, OperationalError

from app.domain.capabilities import require_capabilities
from app.domain.people import normalize_sales_owner_personnel, ensure_unique_personnel_fields
from app.enums import Capability, PersonnelType, EmploymentStatus
from app.models.core import Personnel, Role, Holiday
from app.services import data_backup as backup
from app.services.excel_common import MAX_ROWS, MAX_FILE_BYTES as MAX_FILE_BYTES, fingerprint, read_rows, scalar
from app.services.personnel_master import annual_employment_mm, refresh_personnel_projections, refresh_open_kpi

SHEET = "인력관리"
COLUMNS = [
    ("personnel_id", "id"), ("인력 유형", "personnel_type"), ("사번", "employee_no"), ("성명", "name"),
    ("이메일", "email"), ("본부", "group_name"), ("팀", "team_name"), ("직위", "position_name"),
    ("역할 코드", "role_id"), ("역할", None), ("재직상태", "employment_status"),
    ("MM 시작일", "mm_start_date"), ("MM 종료일", "mm_end_date"), ("연간 재직 MM", None),
    ("사용여부", "is_active"), ("비고", "note"),
]
LABELS = {field: label for label, field in COLUMNS if field}
TYPE_LABELS = {PersonnelType.PMO: "PMO본부 인력", PersonnelType.SALES: "영업대표", PersonnelType.TECHNICAL: "기술본부 인력", PersonnelType.EXTERNAL: "기타/회사 외부 인력"}
STATUS_LABELS = {status: label for label, status in backup.EMPLOYMENT_STATUS_MAP.items()}
VALIDATIONS: dict[str, dict] = {}
TTL = timedelta(minutes=30)
MODELS = (*backup.TABLE_MODELS.values(), Role, Holiday)


def export_workbook(session, user, ids=None):
    require_capabilities(session, user, Capability.PERSONNEL_EXCEL_MANAGE)
    people = {p.id: p for p in session.scalars(select(Personnel)).all()}
    selected = set(people) if ids is None else set(ids)
    if not selected <= people.keys():
        raise HTTPException(status_code=409, detail="다운로드 대상 인력이 변경되었습니다. 다시 조회해 주세요.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append([None] * len(COLUMNS))
    for start, end, title in [(1, 4, "인력 식별/유형"), (5, 10, "기본 정보"), (11, 14, "재직/MM 정보"), (15, 16, "기타")]:
        sheet.cell(1, start, title)
        sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
    sheet.append([label for label, _ in COLUMNS])
    for person in sorted((people[key] for key in selected), key=lambda p: (p.name, p.id)):
        row = []
        for label, field in COLUMNS:
            value = getattr(person, field) if field else None
            if field == "personnel_type":
                value = TYPE_LABELS[person.personnel_type]
            elif field == "role_id":
                value = person.role.code if person.role else None
            elif label == "역할":
                value = person.role.name if person.role else person.role_name
            elif label == "연간 재직 MM":
                value = annual_employment_mm(session, person)
            elif field == "employment_status":
                value = STATUS_LABELS[person.employment_status]
            elif field == "is_active":
                value = "사용" if value else "미사용"
            row.append(scalar(value))
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "E3"
    sheet.auto_filter.ref = f"A2:P{sheet.max_row}"
    for index, (label, _) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(2, index).column_letter].width = 39 if label == "personnel_id" else 23
    for letter in ("A", "C", "E", "I"):
        sheet.column_dimensions[letter].number_format = "@"
    for row in sheet.iter_rows(max_row=2):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EAF0FA")
    raw = BytesIO()
    workbook.save(raw)
    return raw.getvalue()


def validate_workbook(session, user, file_name, raw, personnel_type=None):
    require_capabilities(session, user, Capability.PERSONNEL_EXCEL_MANAGE)
    rows = read_rows(file_name, raw, SHEET)
    required = {label for label, field in COLUMNS if field}
    header_index = next((i for i, row in enumerate(rows[:4]) if required <= {backup._normalize_cell(c.value) for c in row}), None)
    if header_index is None:
        raise HTTPException(status_code=400, detail="personnel_id·인력 유형을 포함한 새 인력 전용 양식이 필요합니다. 템플릿을 다운로드해 주세요.")
    headers = [backup._normalize_cell(c.value) for c in rows[header_index]]
    if len([h for h in headers if h]) != len({h for h in headers if h}):
        raise HTTPException(status_code=400, detail="중복된 컬럼 헤더가 있습니다.")
    original_hash = fingerprint(session, MODELS)
    people = {p.id: p for p in session.scalars(select(Personnel)).all()}
    roles = {r.code: r for r in session.scalars(select(Role)).all()}
    role_codes = {role.id: role.code for role in roles.values()}
    preview, payloads, issues, seen = [], [], [], set()

    def issue(number, label, message):
        issues.append({"row_number": number, "column": label, "message": message})

    def display(field, value):
        if value is None:
            return None
        if field == "personnel_type":
            return TYPE_LABELS[PersonnelType(value)]
        if field == "employment_status":
            return STATUS_LABELS[EmploymentStatus(value)]
        if field == "role_id":
            return role_codes.get(value, value)
        if field == "is_active":
            return "사용" if value else "미사용"
        return scalar(value)

    for number, cells in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(c.value is not None for c in cells):
            continue
        start_issues = len(issues)
        record = {header: c.value for header, c in zip(headers, cells) if header}
        key = backup._normalize_cell(record.get("personnel_id"))
        current = people.get(key)
        if key and (current is None or key in seen):
            issue(number, "personnel_id", "존재하지 않거나 파일 내 중복된 식별자입니다. 신규 행은 비우고 기존 행은 다운로드한 ID를 유지해 주세요.")
        if key:
            seen.add(key)
        values = {}
        for label, field in COLUMNS:
            if not field or field in {"id", "is_active"}:
                continue
            text = backup._normalize_cell(record.get(label))
            try:
                if field == "personnel_type":
                    aliases = {label: kind for kind, label in TYPE_LABELS.items()} | {"PMO": PersonnelType.PMO, "기술": PersonnelType.TECHNICAL, "외부": PersonnelType.EXTERNAL}
                    values[field] = aliases[text] if text in aliases else PersonnelType(text)
                elif field == "employment_status":
                    values[field] = backup.EMPLOYMENT_STATUS_MAP[text] if text in backup.EMPLOYMENT_STATUS_MAP else EmploymentStatus(text) if text else (current.employment_status if current else EmploymentStatus.ACTIVE)
                elif field == "role_id":
                    if text and text not in roles:
                        raise ValueError("역할 코드가 기준정보에 없습니다.")
                    values[field] = roles[text].id if text else None
                elif field.endswith("_date"):
                    value = record.get(label)
                    date_issues = []
                    values[field] = value.date() if isinstance(value, datetime) else backup._parse_date(text, sheet=SHEET, row_number=number, column=label, issues=date_issues)
                    for problem in date_issues:
                        issue(number, label, problem.message)
                else:
                    values[field] = text
                    maximum = getattr(Personnel.__table__.columns[field].type, "length", None)
                    if maximum and text and len(text) > maximum:
                        raise ValueError(f"{maximum}자 이하로 입력해 주세요.")
            except (ValueError, TypeError) as exc:
                issue(number, label, str(exc))
        bool_issues = []
        active = backup._parse_bool(backup._normalize_cell(record.get("사용여부")), sheet=SHEET, row_number=number, column="사용여부", issues=bool_issues)
        values["is_active"] = current.is_active if active is None and current else True if active is None else active
        for problem in bool_issues:
            issue(number, "사용여부", problem.message)
        for header, cell in zip(headers, cells):
            if cell.data_type == "f" and header not in {"역할", "연간 재직 MM"}:
                issue(number, header or "-", "수식 대신 확정된 값을 입력해 주세요.")
        if personnel_type and (current.personnel_type if current else values.get("personnel_type")) != personnel_type:
            issue(number, "인력 유형", "선택한 업로드 보기 범위에 속하지 않는 행입니다. 전체 보기에서 업로드해 주세요.")
        if len(issues) == start_issues:
            try:
                role = normalize_sales_owner_personnel(session, values, current, strict_pmo=True, check_name_uniqueness=False)
                values["role_name"] = current.role_name if current and current.role_id == values.get("role_id") else role.name if role else None
                ensure_unique_personnel_fields(session, employee_no=values.get("employee_no"), email=values.get("email"), personnel_id=key)
                if not values.get("name") or not values.get("group_name"):
                    raise HTTPException(status_code=400, detail="성명과 소속은 필수입니다.")
                if not key and any(p.name == values["name"] and p.group_name == values["group_name"] and p.position_name == values.get("position_name") for p in people.values()):
                    raise HTTPException(status_code=409, detail="같은 성명·소속·직위의 기존 인력이 있습니다. 이름으로 추정 수정하지 않으므로 기존 ID를 입력해 주세요.")
            except HTTPException as exc:
                issue(number, "필수값/식별", str(exc.detail))
        changes = [{"field": LABELS.get(field, "역할"), "before": display(field, getattr(current, field)) if current else None, "after": display(field, value)} for field, value in values.items() if current is None or scalar(getattr(current, field)) != scalar(value)]
        action = "error" if len(issues) > start_issues else "new" if current is None else "update" if changes else "unchanged"
        preview.append({"row_number": number, "code": key or f"신규 {number}행", "name": values.get("name") or "-", "action": action, "changes": changes})
        payloads.append({"id": key, "values": values, "action": action})
        if len(preview) > MAX_ROWS:
            raise HTTPException(status_code=400, detail=f"파일은 {MAX_ROWS}행 이내로 입력해 주세요.")
    if not preview:
        raise HTTPException(status_code=400, detail="업로드할 인력 행이 없습니다.")
    # Validate the complete resulting identity set, including active sales names.
    final = {key: {column.name: getattr(p, column.name) for column in Personnel.__table__.columns} for key, p in people.items()}
    for index, row in enumerate(payloads):
        if row["action"] != "error":
            final[row["id"] or f"new-{index}"] = row["values"]
    for index, row in enumerate(payloads):
        if row["action"] == "error":
            continue
        values, identity = row["values"], row["id"] or f"new-{index}"
        others = [other for key, other in final.items() if key != identity]
        errors = []
        for field in ("employee_no", "email"):
            if values.get(field) and any(p.get(field) == values[field] for p in others):
                errors.append(f"{LABELS[field]} 중복")
        for other in others:
            if values.get("is_active") and other.get("is_active") and values.get("name") == other.get("name") and PersonnelType.SALES in {values.get("personnel_type"), other.get("personnel_type")}:
                errors.append("활성 영업대표 성명이 다른 활성 인력과 중복됩니다.")
                break
        if not row["id"] and any(other.get("name") == values.get("name") and other.get("group_name") == values.get("group_name") and other.get("position_name") == values.get("position_name") for other in others):
            errors.append("신규 인력의 성명·소속·직위가 파일 내 중복됩니다.")
        if errors:
            issue(preview[index]["row_number"], "식별 중복", ", ".join(errors))
            row["action"] = preview[index]["action"] = "error"
    if fingerprint(session, MODELS) != original_hash:
        raise HTTPException(status_code=409, detail="검증 중 DB가 변경되었습니다. 다시 검증해 주세요.")
    now = backup._now()
    for token, cached in list(VALIDATIONS.items()):
        if cached["created_at"] + TTL < now:
            VALIDATIONS.pop(token, None)
    if len(VALIDATIONS) >= 100:
        raise HTTPException(status_code=429, detail="검증 요청이 많습니다. 잠시 후 다시 시도해 주세요.")
    token = str(uuid4())
    result = {"validation_id": token, "file_sha256": backup._file_sha256(raw), "file_name": file_name,
              "summary": {a: sum(r["action"] == a for r in preview) for a in ("new", "update", "unchanged", "error")},
              "rows": preview, "missing": [{"code": p.id, "name": p.name, "is_active": p.is_active} for p in people.values() if p.id not in seen and (not personnel_type or p.personnel_type == personnel_type)],
              "issues": issues, "warnings": ["누락 확인 범위: " + (TYPE_LABELS[personnel_type] if personnel_type else "전체 인력"), "역할 표시명·연간 MM은 입력에서 제외하고 서버가 계산합니다. 영업대표 MM 입력은 기존 정책에 따라 저장하지 않습니다.", "사번·이메일·MM은 원본값을 사용하고, 신규 행은 ID를 비웁니다. 다음 수정은 반영 후 다운로드한 ID를 사용해 주세요."]}
    VALIDATIONS[token] = {"created_at": now, "user_id": user.id, "fingerprint": original_hash, "result": result, "payloads": payloads}
    return result


def apply_workbook(session, user, validation_id, file_sha256, missing_actions):
    require_capabilities(session, user, Capability.PERSONNEL_EXCEL_MANAGE)
    cached = VALIDATIONS.get(validation_id)
    if not cached or cached["created_at"] + TTL < backup._now():
        raise HTTPException(status_code=409, detail="검증 결과가 없거나 만료되었습니다. 다시 검증해 주세요.")
    if cached["user_id"] != user.id:
        raise HTTPException(status_code=403, detail="다른 사용자의 검증 결과는 반영할 수 없습니다.")
    result = cached["result"]
    if result["file_sha256"] != file_sha256 or result["summary"]["error"]:
        raise HTTPException(status_code=409, detail="검증 파일이 다르거나 오류가 있습니다.")
    if not set(missing_actions) <= {r["code"] for r in result["missing"]} or any(v not in {"keep", "deactivate"} for v in missing_actions.values()):
        raise HTTPException(status_code=400, detail="검증된 누락 인력의 유지/비활성화만 선택할 수 있습니다.")
    backup._acquire_lock(session, locked_by=user.name)
    counts = {"new": 0, "update": 0, "unchanged": 0, "deactivated": 0}
    created_ids = []
    projections_changed = False
    before = None

    def record_failure(message):
        try:
            backup._append_manifest({"id": backup._make_id("OPR"), "created_at": backup._now().isoformat(timespec="seconds"), "kind": "upload_personnel_failed", "status": "failed", "actor_name": user.name, "backup_id": None, "source_backup_id": before["backup_id"] if before else None, "source_file_name": result["file_name"], "message": message})
        except OSError:
            pass  # Preserve the original failure if the audit disk also fails.
    try:
        if session.get_bind().dialect.name == "sqlite":
            session.connection().exec_driver_sql("BEGIN IMMEDIATE")
        else:
            for model in MODELS:
                session.scalars(select(model).order_by(model.id).with_for_update()).all()
        require_capabilities(session, user, Capability.PERSONNEL_EXCEL_MANAGE)
        if VALIDATIONS.get(validation_id) is not cached or fingerprint(session, MODELS) != cached["fingerprint"]:
            raise HTTPException(status_code=409, detail="검증 후 DB가 변경되었습니다. 다시 검증해 주세요.")
        before = backup.create_backup(session, actor_name=user.name, kind="personnel_upload_before_backup", source_file_name=result["file_name"])
        for row in cached["payloads"]:
            counts[row["action"]] += 1
            if row["action"] == "unchanged":
                continue
            person = session.get(Personnel, row["id"]) if row["id"] else Personnel()
            needs_projection = row["action"] == "new" or any(
                getattr(person, field) != value for field, value in row["values"].items()
                if field in {"personnel_type", "mm_start_date", "mm_end_date", "employment_status", "is_active"}
            )
            for field, value in row["values"].items():
                setattr(person, field, value)
            session.add(person)
            session.flush()
            if needs_projection:
                refresh_personnel_projections(session, person)
                projections_changed = True
            if row["action"] == "new":
                created_ids.append({"id": person.id, "name": person.name})
        for key, action in missing_actions.items():
            person = session.get(Personnel, key)
            if action == "deactivate" and person.is_active:
                person.is_active = False
                refresh_personnel_projections(session, person)
                projections_changed = True
                counts["deactivated"] += 1
        if projections_changed:
            refresh_open_kpi(session)
        session.commit()
    except (IntegrityError, OperationalError) as exc:
        session.rollback()
        record_failure("DB 충돌로 전체 반영 롤백")
        raise HTTPException(status_code=409, detail="인력 반영 중 충돌이 발생해 롤백했습니다. 다시 검증해 주세요.") from exc
    except Exception as exc:
        session.rollback()
        record_failure(str(exc))
        raise
    finally:
        backup._release_lock(session)
    VALIDATIONS.pop(validation_id, None)
    warning = None
    try:
        backup._append_manifest({"id": backup._make_id("OPR"), "created_at": backup._now().isoformat(timespec="seconds"), "kind": "upload_personnel_complete", "status": "success", "actor_name": user.name, "backup_id": None, "source_backup_id": before["backup_id"], "source_file_name": result["file_name"], "counts": counts})
    except OSError:
        warning = "DB 반영은 완료했으나 작업 이력 파일 저장에 실패했습니다. 업로드 전 백업은 보존됩니다."
    return {"counts": counts, "created_ids": created_ids, "pre_backup_id": before["backup_id"], "message": warning or "인력 반영을 완료했습니다. 식별자가 포함된 파일을 다시 다운로드해 주세요."}
