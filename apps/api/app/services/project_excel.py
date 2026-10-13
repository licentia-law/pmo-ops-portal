"""Project-only workbook validation and transactional master import."""
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session

from app.domain.capabilities import require_capabilities
from app.domain.projects import get_missing_project_fields, is_valid_status_transition, project_team_members
from app.enums import Capability, ProjectLogStatus, ProjectStatus, ProjectType
from app.models.core import CurrentAssignmentSnapshot, Project, ProjectCode, ProjectLog, User, WeeklyLoadSnapshot
from app.schemas.projects import ProjectMasterCreate, ProjectMasterUpdate
from app.services import data_backup as backup
from app.services.excel_common import fingerprint
from app.services.project_master import (
    PROJECT_MASTER_SYNC_FIELDS, ProjectMasterConflictError, ProjectMasterValidationError,
    create_project_master, update_project_master,
)

SHEET = "프로젝트관리"
# Both download and upload use this business-column contract.
COLUMNS = [
    ("코드", "code"), ("사업명", "name"), ("고객사", "client_name"),
    ("사업유형", "project_type"), ("확도", "certainty"),
    ("총 사업금액", "total_amount"), ("당사 사업금액", "company_amount"),
    ("공고번호", "bid_notice_no"), ("공고일", "bid_notice_date"), ("상태", "status"),
    ("영업부서", "sales_department"), ("영업대표", "sales_owner"),
    ("제안PM", "proposal_pm_name"), ("발표PM", "presentation_pm_name"), ("수행PM", "delivery_pm_name"),
    ("제안/수행팀", None), ("사업 시작일", "start_date"), ("사업 종료일", "end_date"),
    ("제안서 제출일", "submission_at"), ("제출 형식", "submission_format"), ("제출 유의사항", "submission_note"),
    ("제안 발표일", "presentation_at"), ("발표 형식", "presentation_format"), ("발표 유의사항", "presentation_note"),
    ("최근 활동일", None), ("사용여부", "is_active"),
]
LABELS = {field: label for label, field in COLUMNS if field}
TYPE_LABELS = {value: key for key, value in backup.PROJECT_TYPE_MAP.items()}
STATUS_LABELS = {value: key for key, value in backup.PROJECT_STATUS_MAP.items()}
VALIDATIONS: dict[str, dict] = {}
TTL = timedelta(minutes=30)
MAX_ROWS = 5000
MAX_FILE_BYTES = 20 * 1024 * 1024


def _scalar(value):
    if isinstance(value, datetime):
        return value.isoformat()
    return backup._serialize_scalar(value)


def _fingerprint(session: Session) -> str:
    return fingerprint(session, (ProjectCode, Project))


def _projects_by_code(session: Session) -> dict[str, Project]:
    rows = session.scalars(select(Project)).all()
    codes = session.scalars(select(ProjectCode)).all()
    linked = {p.project_code_id for p in rows}
    if any(c.id not in linked for c in codes) or any(
        p.project_code is None or any(getattr(p, field) != getattr(p.project_code, field) for field in PROJECT_MASTER_SYNC_FIELDS) for p in rows
    ):
        raise HTTPException(status_code=409, detail="프로젝트와 코드의 연결 정합성을 먼저 확인해 주세요.")
    return {p.code: p for p in rows}


def _display(value):
    if isinstance(value, bool):
        return "사용" if value else "미사용"
    return _scalar(value)


def export_workbook(session: Session, user: User, codes: list[str] | None) -> bytes:
    require_capabilities(session, user, Capability.PROJECT_EXCEL_MANAGE)
    projects = _projects_by_code(session)
    selected = set(projects) if codes is None else set(codes)
    if not selected <= set(projects):
        raise HTTPException(status_code=409, detail="다운로드 대상 프로젝트가 변경되었습니다. 다시 조회해 주세요.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append([None] * len(COLUMNS))
    for start, end, title in [(1, 2, "프로젝트 코드/사업명"), (3, 10, "프로젝트 기본정보"), (11, 16, "인력 정보"), (17, 24, "프로젝트 일정"), (25, 26, "기타")]:
        sheet.cell(1, start, title)
        sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
    sheet.append([label for label, _ in COLUMNS])
    teams = project_team_members(session, [projects[code].id for code in selected])
    recent = dict(session.execute(select(ProjectLog.project_id, func.max(ProjectLog.logged_at)).group_by(ProjectLog.project_id)).all())
    for code in sorted(selected, reverse=True):
        project = projects[code]
        row = []
        excluded = {project.proposal_pm_name, project.presentation_pm_name, project.delivery_pm_name}
        for label, field in COLUMNS:
            if label == "제안/수행팀":
                value = " / ".join(member for member in teams.get(project.id, []) if member.split(" ", 1)[0] not in excluded)
            elif label == "최근 활동일":
                value = recent.get(project.id)
            elif field == "is_active":
                value = "사용" if project.project_code.is_active else "미사용"
            elif field == "project_type":
                value = TYPE_LABELS[project.project_type]
            elif field == "status":
                value = STATUS_LABELS[project.status]
            else:
                value = getattr(project, field)
            if isinstance(value, datetime):
                value = value.isoformat(sep=" ")
            elif isinstance(value, date):
                value = value.isoformat()
            elif isinstance(value, Decimal):
                value = float(value)
            row.append(value)
        sheet.append(row)
        # Strings are literal data, including names starting with '='.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "C3"
    sheet.auto_filter.ref = f"A2:Z{max(sheet.max_row, 2)}"
    for cells in sheet.iter_rows():
        for cell in cells:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.row <= 2:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DEE3EB" if cell.row == 1 else "EAF0FA")
    for index, (label, _) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(2, index).column_letter].width = 30 if label in {"사업명", "제안/수행팀", "제출 유의사항", "발표 유의사항"} else 19
    raw = BytesIO()
    workbook.save(raw)
    return raw.getvalue()


def _number(value) -> float | None:
    text = backup._normalize_cell(value)
    if text is None:
        return None
    try:
        number = Decimal(text.replace(",", ""))
        if not number.is_finite() or number < 0 or number >= Decimal("100000000000000") or number != number.quantize(Decimal("0.01")):
            raise ValueError("금액은 0 이상, 소수 둘째 자리까지 입력해 주세요.")
        return float(number)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("금액은 0 이상, 소수 둘째 자리까지의 숫자로 입력해 주세요.") from exc


def _datetime_cell(value, issues, row_number, label):
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            issues.append(backup.ValidationIssue(SHEET, row_number, label, str(value), "error", "시간대가 없는 현지 날짜/시간을 입력해 주세요."))
            return None
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = backup._normalize_cell(value)
    if not text:
        return None
    try:
        result = datetime.fromisoformat(text)
        if result.tzinfo is not None:
            raise ValueError("시간대가 없는 현지 날짜/시간을 입력해 주세요.")
        return result
    except ValueError:
        return backup._parse_datetime(text, sheet=SHEET, row_number=row_number, column=label, issues=issues)


def validate_workbook(session: Session, user: User, file_name: str, raw: bytes) -> dict:
    require_capabilities(session, user, Capability.PROJECT_EXCEL_MANAGE)
    if not file_name.lower().endswith(".xlsx") or not raw or len(raw) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="20MB 이하의 .xlsx 파일을 선택해 주세요.")
    try:
        workbook = load_workbook(BytesIO(raw), data_only=False, read_only=True)
        if workbook.sheetnames != [SHEET]:
            raise ValueError("프로젝트관리 시트만 포함된 프로젝트 전용 파일이 필요합니다.")
        sheet = workbook[SHEET]
        if sheet.max_row > MAX_ROWS + 4 or sheet.max_column > 60:
            raise ValueError(f"프로젝트 파일은 {MAX_ROWS}행 이내로 입력해 주세요.")
        rows = list(sheet.iter_rows())
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"프로젝트 엑셀을 읽을 수 없습니다: {exc}") from exc
    finally:
        if "workbook" in locals():
            workbook.close()
    required = {label for label, field in COLUMNS if field and field not in {"total_amount", "company_amount"}}
    header_index = next((i for i, row in enumerate(rows[:4]) if required <= {backup._normalize_cell(cell.value) for cell in row}), None)
    if header_index is None:
        raise HTTPException(status_code=400, detail="프로젝트 목록의 필수 컬럼 헤더를 찾을 수 없습니다.")
    headers = [backup._normalize_cell(cell.value) for cell in rows[header_index]]
    if len([h for h in headers if h]) != len({h for h in headers if h}):
        raise HTTPException(status_code=400, detail="중복된 컬럼 헤더가 있습니다.")
    legacy = "총 사업금액" not in headers or "당사 사업금액" not in headers
    if legacy and ("사업금액" not in headers or "총 사업금액" in headers or "당사 사업금액" in headers):
        raise HTTPException(status_code=400, detail="총 사업금액과 당사 사업금액 컬럼이 필요합니다.")
    fingerprint = _fingerprint(session)
    projects = _projects_by_code(session)
    seen = set()
    preview = []
    payloads = []
    all_issues = []
    for row_number, cells in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(cell.value is not None for cell in cells):
            continue
        record = {header: cell.value for header, cell in zip(headers, cells) if header}
        issues = []
        code = backup._normalize_cell(record.get("코드"))
        project = projects.get(code)
        values = {}
        for label, field in COLUMNS:
            if not field or field in {"total_amount", "company_amount", "is_active"}:
                continue
            value = record.get(label)
            try:
                if field in {"project_type", "status"}:
                    enum = ProjectType if field == "project_type" else ProjectStatus
                    mapping = backup.PROJECT_TYPE_MAP if field == "project_type" else backup.PROJECT_STATUS_MAP
                    text = backup._normalize_cell(value)
                    values[field] = mapping[text] if text in mapping else enum(text)
                elif field.endswith("_date"):
                    values[field] = value.date() if isinstance(value, datetime) else backup._parse_date(backup._normalize_cell(value), sheet=SHEET, row_number=row_number, column=label, issues=issues)
                elif field.endswith("_at"):
                    values[field] = _datetime_cell(value, issues, row_number, label)
                else:
                    values[field] = backup._normalize_cell(value)
                    column = Project.__table__.columns[field]
                    maximum = getattr(column.type, "length", None)
                    if maximum and values[field] and len(values[field]) > maximum:
                        raise ValueError(f"{maximum}자 이하로 입력해 주세요.")
            except (ValueError, TypeError) as exc:
                issues.append(backup.ValidationIssue(SHEET, row_number, label, str(value), "error", str(exc)))
        if not code or not re.fullmatch(r"P[0-9]{7}", code) or code in seen:
            issues.append(backup.ValidationIssue(SHEET, row_number, "코드", code or "-", "error", "코드는 P+연도+3자리 번호 형식이며 파일 내 중복될 수 없습니다."))
        if code:
            seen.add(code)
        for cell, header in zip(cells, headers):
            if cell.data_type == "f" and header not in {"제안/수행팀", "최근 활동일"}:
                issues.append(backup.ValidationIssue(SHEET, row_number, header or "-", str(cell.value), "error", "수식 대신 확정된 값을 입력해 주세요."))
        if legacy:
            total, company, _ = backup._parse_amount_text(backup._normalize_cell(record.get("사업금액")), row_number=row_number, issues=issues)
            values.update(total_amount=total, company_amount=company)
            for field in ("total_amount", "company_amount"):
                try:
                    values[field] = _number(values[field])
                except ValueError as exc:
                    values[field] = None
                    issues.append(backup.ValidationIssue(SHEET, row_number, "사업금액", str(record.get("사업금액")), "error", str(exc)))
            # The old download rounded the displayed amount; unchanged display text
            # must retain the existing precise numbers rather than import that rounding.
            if project and backup._normalize_cell(record.get("사업금액")) == backup._normalize_cell(project.amount_text):
                values.update(total_amount=project.total_amount, company_amount=project.company_amount)
        else:
            for field in ("total_amount", "company_amount"):
                try:
                    values[field] = _number(record.get(LABELS[field]))
                except ValueError as exc:
                    issues.append(backup.ValidationIssue(SHEET, row_number, LABELS[field], str(record.get(LABELS[field])), "error", str(exc)))
        active = backup._parse_bool(backup._normalize_cell(record.get("사용여부")), sheet=SHEET, row_number=row_number, column="사용여부", issues=issues)
        active = project.project_code.is_active if active is None and project else (True if active is None else active)
        missing = get_missing_project_fields({**values, "project_code_id": project.project_code_id if project else "pending"})
        if missing:
            issues.append(backup.ValidationIssue(SHEET, row_number, "필수값", "-", "error", f"필수 항목 누락: {', '.join(missing)}"))
        if values.get("start_date") and values.get("end_date") and values["start_date"] > values["end_date"]:
            issues.append(backup.ValidationIssue(SHEET, row_number, "사업 종료일", "-", "error", "종료일은 시작일보다 빠를 수 없습니다."))
        if values.get("total_amount") is not None and values.get("company_amount") is not None and values["company_amount"] > values["total_amount"]:
            issues.append(backup.ValidationIssue(SHEET, row_number, "당사 사업금액", "-", "error", "당사 금액은 총 금액을 초과할 수 없습니다."))
        if project and values.get("status") is not None and not is_valid_status_transition(project.status, values["status"]):
            issues.append(backup.ValidationIssue(SHEET, row_number, "상태", "-", "error", "허용되지 않는 상태 전환입니다."))
        changes = []
        if project:
            for field, value in values.items():
                if _scalar(getattr(project, field)) != _scalar(value):
                    changes.append({"field": LABELS[field], "before": _display(getattr(project, field)), "after": _display(value)})
            if project.project_code.is_active != active:
                changes.append({"field": "사용여부", "before": _display(project.project_code.is_active), "after": _display(active)})
        else:
            changes = [{"field": LABELS[field], "before": None, "after": _display(value)} for field, value in values.items() if value is not None]
        action = "error" if issues else ("new" if project is None else ("update" if changes else "unchanged"))
        preview.append({"row_number": row_number, "code": code or "-", "name": values.get("name") or "-", "action": action, "changes": changes})
        if len(preview) > MAX_ROWS:
            raise HTTPException(status_code=400, detail=f"프로젝트 파일은 {MAX_ROWS}행 이내로 입력해 주세요.")
        payloads.append({"values": values, "is_active": active, "action": action})
        all_issues.extend(issue.to_dict() for issue in issues)
    if not preview:
        raise HTTPException(status_code=400, detail="업로드할 프로젝트 행이 없습니다.")
    if _fingerprint(session) != fingerprint:
        raise HTTPException(status_code=409, detail="검증 중 프로젝트 데이터가 변경되었습니다. 다시 검증해 주세요.")
    missing_rows = [{"code": p.code, "name": p.name, "is_active": p.project_code.is_active} for p in projects.values() if p.code not in seen]
    now = backup._now()
    for token, cached in list(VALIDATIONS.items()):
        if cached["created_at"] + TTL < now:
            VALIDATIONS.pop(token, None)
    if len(VALIDATIONS) >= 100:
        raise HTTPException(status_code=429, detail="검증 요청이 많습니다. 잠시 후 다시 시도해 주세요.")
    token = str(uuid4())
    digest = backup._file_sha256(raw)
    result = {
        "validation_id": token, "file_sha256": digest, "file_name": file_name,
        "summary": {action: sum(row["action"] == action for row in preview) for action in ("new", "update", "unchanged", "error")},
        "rows": preview, "missing": sorted(missing_rows, key=lambda row: row["code"]), "issues": all_issues,
        "warnings": ["제안/수행팀과 최근 활동일은 참고용이며 업로드 입력에서 제외됩니다."] + (["기존 사업금액 열을 호환 형식으로 읽었습니다. 기존 표시와 동일하면 정확한 저장 금액을 유지합니다. 금액 수정은 분리 금액 양식을 사용해 주세요."] if legacy else []),
    }
    VALIDATIONS[token] = {"created_at": now, "user_id": user.id, "file_sha256": digest, "fingerprint": fingerprint, "payloads": payloads, "result": result}
    return result


def _refresh_current_names(session: Session, project: Project) -> None:
    today = date.today()
    for row in session.scalars(select(CurrentAssignmentSnapshot).where(CurrentAssignmentSnapshot.as_of_date == today, CurrentAssignmentSnapshot.project_id == project.id)).all():
        row.project_name, row.project_code = project.name, project.code
    for row in session.scalars(select(CurrentAssignmentSnapshot).where(CurrentAssignmentSnapshot.as_of_date == today, CurrentAssignmentSnapshot.next_project_id == project.id)).all():
        row.next_project_name = project.name
    for row in session.scalars(select(WeeklyLoadSnapshot).where(WeeklyLoadSnapshot.as_of_date == today, WeeklyLoadSnapshot.project_id == project.id)).all():
        row.project_name = project.name


def apply_workbook(session: Session, user: User, validation_id: str, file_sha256: str, missing_actions: dict[str, str]) -> dict:
    require_capabilities(session, user, Capability.PROJECT_EXCEL_MANAGE)
    cached = VALIDATIONS.get(validation_id)
    if cached is None or cached["created_at"] + TTL < backup._now():
        raise HTTPException(status_code=409, detail="검증 결과가 없거나 만료되었습니다. 다시 검증해 주세요.")
    if cached["user_id"] != user.id:
        raise HTTPException(status_code=403, detail="다른 사용자의 검증 결과는 반영할 수 없습니다.")
    if file_sha256 != cached["file_sha256"] or cached["result"]["summary"]["error"]:
        raise HTTPException(status_code=409, detail="검증한 파일과 다르거나 검증 오류가 있습니다.")
    missing = {row["code"] for row in cached["result"]["missing"]}
    if not set(missing_actions) <= missing or any(action not in {"keep", "deactivate"} for action in missing_actions.values()):
        raise HTTPException(status_code=400, detail="검증된 누락 프로젝트의 유지/비활성화만 선택할 수 있습니다.")
    backup._acquire_lock(session, locked_by=user.name)
    before_backup = None
    try:
        if session.get_bind().dialect.name == "sqlite":
            session.connection().exec_driver_sql("BEGIN IMMEDIATE")
        else:
            session.scalars(select(ProjectCode).order_by(ProjectCode.id).with_for_update()).all()
            session.scalars(select(Project).order_by(Project.id).with_for_update()).all()
        require_capabilities(session, user, Capability.PROJECT_EXCEL_MANAGE)
        if VALIDATIONS.get(validation_id) is not cached or _fingerprint(session) != cached["fingerprint"]:
            raise HTTPException(status_code=409, detail="검증 후 프로젝트 데이터가 변경되었습니다. 다시 검증해 주세요.")
        before_backup = backup.create_backup(session, actor_name=user.name, kind="project_upload_before_backup", source_file_name=cached["result"]["file_name"])
        projects = _projects_by_code(session)
        counts = {"new": 0, "update": 0, "unchanged": 0, "deactivated": 0}
        for row in cached["payloads"]:
            action = row["action"]
            counts[action] += 1
            if action == "unchanged":
                continue
            values = dict(row["values"])
            values["amount_text"] = f"{backup._format_amount_value(values['total_amount'])}/{backup._format_amount_value(values['company_amount'])}"
            code_values = {field: values[field] for field in ("code", "name", "project_type", "status", "certainty")}
            code_values["is_active"] = row["is_active"]
            if action == "new":
                _, project = create_project_master(session, user, ProjectMasterCreate(project_code=code_values, project=values))
            else:
                project = projects[values["code"]]
                previous_status = project.status
                update_project_master(session, user, project, ProjectMasterUpdate(project_code=code_values, project=values))
                if project.status == previous_status:
                    session.add(ProjectLog(project_id=project.id, log_status=ProjectLogStatus.MEMO, logged_at=backup._now(), author_name=user.name, updated_by_name=user.name, content="프로젝트 엑셀 수정"))
            session.flush()
            project.recent_activity_at = session.scalar(select(func.max(ProjectLog.logged_at)).where(ProjectLog.project_id == project.id))
            _refresh_current_names(session, project)
        for code, action in missing_actions.items():
            if action == "deactivate" and projects[code].project_code.is_active:
                project = projects[code]
                project.project_code.is_active = False
                session.add(ProjectLog(project_id=project.id, log_status=ProjectLogStatus.MEMO, logged_at=backup._now(), author_name=user.name, updated_by_name=user.name, content="프로젝트 엑셀 누락: 비활성화"))
                session.flush()
                project.recent_activity_at = session.scalar(select(func.max(ProjectLog.logged_at)).where(ProjectLog.project_id == project.id))
                counts["deactivated"] += 1
        session.commit()
    except (ProjectMasterValidationError, ProjectMasterConflictError, IntegrityError, OperationalError) as exc:
        session.rollback()
        raise HTTPException(status_code=409, detail="프로젝트 반영 중 충돌이 발생했습니다. 데이터가 롤백되었으니 다시 검증해 주세요.") from exc
    except Exception:
        session.rollback()
        raise
    finally:
        backup._release_lock(session)
    VALIDATIONS.pop(validation_id, None)
    return {"counts": counts, "pre_backup_id": before_backup["backup_id"], "message": "프로젝트 반영을 완료했습니다."}
