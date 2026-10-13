from datetime import date, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text

from app.core.config import settings
from app.models.core import Personnel, ProjectAssignment, MonthlyAssignmentMM, MonthlyEmploymentMM, CurrentAssignmentSnapshot, JobLock, User
from app.services import assignment_import as importer, data_backup as backup
from test_capabilities import capability_db as capability_db, headers, grant
from test_project_excel import seed, dataset


@pytest.fixture
def conversion(capability_db, monkeypatch):
    client, factory = capability_db
    monkeypatch.setattr(settings, "assignment_import_enabled", True)
    monkeypatch.setattr(importer, "VALIDATIONS", {})
    projects = seed(factory, 2)
    people = {}
    with factory() as session:
        for kind in ("pmo", "sales", "technical", "external"):
            key = str(uuid4())
            people[kind] = key
            session.add(Personnel(id=key, name=f"{kind} 인력", personnel_type=kind, position_name="책임", employment_status="active", is_active=True, mm_start_date=date(date.today().year, 1, 1)))
        session.commit()
    return client, factory, projects, people


def row(people, kind="pmo", **edits):
    today = date.today()
    return {"배정키": str(uuid4()), "프로젝트 코드": "P2026001", "personnel_id": people[kind], "배정 유형": "수행", "프로젝트 역할": "수행팀", "배정 상태": "종료", "시작일": date(today.year, 1, 1).isoformat(), "종료일": date(today.year, 1, 31).isoformat(), "월 기준 MM": 1, "확도율": 50, "상주 유형": "비상주", "대표 여부": "거짓", "표시 순서": 0, "비고": "원본 메모", **edits}


def workbook(rows):
    book = Workbook()
    book.active.title = importer.SHEET
    book.active.append([label for label, _ in importer.COLUMNS])
    for values in rows:
        book.active.append([values.get(label) for label, _ in importer.COLUMNS])
    raw = BytesIO()
    book.save(raw)
    return raw.getvalue()


def validate(client, rows, actor="developer"):
    response = client.post("/api/assignment-import/validate", headers=headers(actor, **{"x-upload-filename": "legacy.xlsx"}), content=workbook(rows))
    assert response.status_code == 200, response.text
    return response.json()["data"]


def apply(client, preview, actor="developer", **changes):
    return client.post("/api/assignment-import/apply", headers=headers(actor), json={"validation_id": preview["validation_id"], "file_sha256": preview["file_sha256"], "acknowledge_warnings": True, **changes})


def test_disabled_production_and_capability(conversion, monkeypatch):
    client, factory, _, people = conversion
    before = dataset(factory)
    for path in ("status", "template", "references"):
        assert client.get(f"/api/assignment-import/{path}", headers=headers("admin")).status_code == 403
    monkeypatch.setattr(settings, "assignment_import_enabled", False)
    assert not client.get("/api/assignment-import/status", headers=headers()).json()["data"]["enabled"]
    assert client.post("/api/assignment-import/validate", headers=headers(), content=workbook([row(people)])).status_code == 403
    grant(client, "member", ["assignment_manage"])
    monkeypatch.setattr(settings, "assignment_import_enabled", True)
    assert client.get("/api/assignment-import/status", headers=headers("member")).status_code == 200
    monkeypatch.setattr(settings, "app_environment", "production")
    with factory() as session:
        assert not importer.status(session, session.get(User, "developer"))["enabled"]
    assert dataset(factory) == before


def test_template_references_and_validation_readonly(conversion):
    client, factory, _, people = conversion
    before = dataset(factory)
    response = client.get("/api/assignment-import/template", headers=headers())
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == [importer.SHEET]
    assert [c.value for c in book.active[1]] == [label for label, _ in importer.COLUMNS]
    refs = load_workbook(BytesIO(client.get("/api/assignment-import/references", headers=headers()).content))
    assert refs.sheetnames == ["프로젝트 참조", "인력 참조", "입력 안내"]
    assert people["pmo"] in [r[0] for r in refs["인력 참조"].iter_rows(min_row=2, values_only=True)]
    preview = validate(client, [row(people)])
    assert preview["summary"] == {"new": 1, "error": 0}
    assert dataset(factory) == before


@pytest.mark.parametrize("edit", [{"personnel_id": "unknown"}, {"personnel_id": str(uuid4())}, {"프로젝트 코드": "missing"}, {"배정키": None}, {"월 기준 MM": None}, {"월 기준 MM": -1}, {"월 기준 MM": 1.555}, {"월 기준 MM": "NaN"}, {"확도율": 101}, {"시작일": "2026-02-01", "종료일": "2026-01-31"}, {"프로젝트 역할": "제안PM"}, {"배정 상태": "invalid"}, {"대표 여부": "maybe"}, {"표시 순서": 1.5}, {"비고": "=1+1"}])
def test_bad_input_no_partial_apply(conversion, edit):
    client, factory, _, people = conversion
    before = dataset(factory)
    preview = validate(client, [row(people, **edit), row(people, "sales")])
    assert preview["summary"]["error"] == 1
    assert apply(client, preview).status_code == 409
    assert dataset(factory) == before


def test_duplicate_and_overlap(conversion):
    client, _, _, people = conversion
    original = row(people)
    assert validate(client, [original, original])["summary"]["error"] == 1
    overlap = row(people, **{"프로젝트 코드": "P2026002"})
    preview = validate(client, [original, overlap])
    assert preview["summary"]["error"] == 0 and any("같은 파일" in w for w in preview["warnings"])
    assert apply(client, preview, acknowledge_warnings=False).status_code == 409
    assert apply(client, preview).status_code == 200


def test_all_types_history_kpi_and_repeat_guard(conversion, monkeypatch):
    client, factory, _, people = conversion
    old = dataset(factory)
    preview = validate(client, [row(people, kind) for kind in people])
    response = apply(client, preview)
    assert response.status_code == 200, response.text
    result = response.json()["data"]
    assert result["counts"]["assignments_created"] == 4 and result["pre_backup_id"]
    after = dataset(factory)
    # Masters, historical MM, snapshots, KPI and assignments keep original IDs/values.
    for table in ("personnel", "monthly_assignment_mm", "monthly_employment_mm", "monthly_kpi_summaries", "current_assignment_snapshots", "weekly_load_snapshots", "project_assignments"):
        indexed = {r["id"]: r for r in after[table]}
        for entry in old[table]:
            if entry.get("as_of_date") == date.today().isoformat():
                continue  # Today's operational projections are deliberately refreshed.
            assert indexed[entry["id"]] == entry
    with factory() as session:
        assert session.execute(text("PRAGMA foreign_key_check")).all() == []
        rows = session.scalars(select(MonthlyAssignmentMM).where(MonthlyAssignmentMM.year == date.today().year, MonthlyAssignmentMM.month == 1, MonthlyAssignmentMM.personnel_id.in_(list(people.values())))).all()
        assert len(rows) == 4 and all(float(r.assignment_mm) == 1 and float(r.weighted_mm) == .5 for r in rows)
        assert not session.scalars(select(MonthlyEmploymentMM).where(MonthlyEmploymentMM.personnel_id.in_([people[k] for k in people if k != "pmo"]))).all()
        current = session.scalars(select(CurrentAssignmentSnapshot).where(CurrentAssignmentSnapshot.as_of_date == date.today())).all()
        assert all(r.personnel_id == people["pmo"] for r in current)
        from app.domain.kpi import pmo_monthly_kpi_values
        assert pmo_monthly_kpi_values(session, date.today().year, 1)["running_mm"] == .5
        marker = importer.completed(session)
        assert client.patch(f"/api/project-logs/{marker.id}", headers=headers(), json={"content": "erase"}).status_code == 403
    monkeypatch.setattr(importer, "VALIDATIONS", {})
    assert client.get("/api/assignment-import/status", headers=headers()).json()["data"]["completed"]
    assert apply(client, preview).status_code == 409
    assert client.post("/api/assignment-import/validate", headers=headers(), content=workbook([row(people)])).status_code == 409


def test_tokens_staleness_and_expiry(conversion):
    client, factory, _, people = conversion
    preview = validate(client, [row(people)])
    assert apply(client, preview, file_sha256="wrong").status_code == 409
    grant(client, "member", ["assignment_manage"])
    assert apply(client, preview, actor="member").status_code == 403
    importer.VALIDATIONS[preview["validation_id"]]["created_at"] -= timedelta(hours=1)
    assert apply(client, preview).status_code == 409
    preview = validate(client, [row(people)])
    with factory() as session:
        session.get(Personnel, people["pmo"]).note = "changed after validation"
        session.commit()
    assert apply(client, preview).status_code == 409
    assert not client.get("/api/assignment-import/status", headers=headers()).json()["data"]["completed"]


def test_cancelled_history_and_retired_history(conversion):
    client, factory, _, people = conversion
    with factory() as session:
        p = session.get(Personnel, people["pmo"])
        p.is_active = False
        p.employment_status = "retired"
        p.mm_end_date = date(date.today().year, 1, 31)
        session.commit()
    preview = validate(client, [row(people), row(people, "external", **{"배정 상태": "취소"})])
    response = apply(client, preview)
    assert response.status_code == 200, response.text
    with factory() as session:
        for kind, expected in (("pmo", 1), ("external", 0)):
            ledger = session.scalar(select(MonthlyAssignmentMM).where(MonthlyAssignmentMM.personnel_id == people[kind], MonthlyAssignmentMM.month == 1))
            assert float(ledger.assignment_mm) == expected


def test_failed_batch_rolls_back_and_retries(conversion, monkeypatch):
    client, factory, _, people = conversion
    before = dataset(factory)
    preview = validate(client, [row(people), row(people, "sales")])
    original = importer.refresh_assignment_mm
    calls = []
    def failing(*args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            raise RuntimeError("simulated second row failure")
        return original(*args, **kwargs)
    monkeypatch.setattr(importer, "refresh_assignment_mm", failing)
    with pytest.raises(RuntimeError, match="second row"):
        apply(client, preview)
    assert dataset(factory) == before
    with factory() as session:
        assert not importer.completed(session)
        assert not session.scalars(select(JobLock)).all()
    assert list(backup.BACKUP_DIR.glob("*.json"))
    monkeypatch.setattr(importer, "refresh_assignment_mm", original)
    assert apply(client, preview).status_code == 200


def test_existing_duplicate_blocks_and_backup_failure(conversion, monkeypatch):
    client, factory, projects, people = conversion
    with factory() as session:
        session.add(ProjectAssignment(project_id=projects[0], personnel_id=people["pmo"], assignment_type="delivery", assignment_role="delivery_team", assignment_status="cancelled", start_date=date(date.today().year, 1, 1), end_date=date(date.today().year, 1, 31)))
        session.commit()
    assert validate(client, [row(people)])["summary"]["error"] == 1
    preview = validate(client, [row(people, "external")])
    before = dataset(factory)
    def fail(*args, **kwargs):
        raise OSError("backup disk failure")
    monkeypatch.setattr(backup, "create_backup", fail)
    with pytest.raises(OSError, match="backup disk"):
        apply(client, preview)
    assert dataset(factory) == before


def test_backup_roundtrip_completion_and_pre_import_recovery(conversion):
    client, factory, _, people = conversion
    before = dataset(factory)
    preview = validate(client, [row(people, kind) for kind in people])
    result = apply(client, preview).json()["data"]
    converted = dataset(factory)
    with factory() as session:
        user = session.get(User, "developer")
        post = backup.create_backup(session, actor_name=user.name, kind="manual")
        backup.restore_backup(session, backup_id=post["backup_id"], memo=None, user=user)
        assert importer.status(session, user)["completed"]
    assert dataset(factory) == converted
    with factory() as session:
        backup.restore_backup(session, backup_id=result["pre_backup_id"], memo=None, user=session.get(User, "developer"))
        assert not importer.completed(session)
    assert dataset(factory) == before
    assert apply(client, validate(client, [row(people)])).status_code == 200


def test_missing_pmo_mm_start_and_period_review_warning(conversion):
    client, factory, _, people = conversion
    with factory() as session:
        session.get(Personnel, people["pmo"]).mm_start_date = None
        session.commit()
    assert validate(client, [row(people)])["summary"]["error"] == 1
    with factory() as session:
        session.get(Personnel, people["pmo"]).mm_start_date = date(date.today().year, 2, 1)
        session.commit()
    preview = validate(client, [row(people)])
    assert any("재직 기간 밖" in warning for warning in preview["warnings"])


def test_employment_history_before_first_assignment_is_filled(conversion):
    client, factory, _, people = conversion
    previous_year = date.today().year - 1
    with factory() as session:
        session.get(Personnel, people["pmo"]).mm_start_date = date(previous_year, 1, 1)
        session.commit()
    assert apply(client, validate(client, [row(people)])).status_code == 200
    with factory() as session:
        ledger = session.scalar(select(MonthlyEmploymentMM).where(MonthlyEmploymentMM.personnel_id == people["pmo"], MonthlyEmploymentMM.year == previous_year, MonthlyEmploymentMM.month == 1))
        assert ledger is not None and float(ledger.employment_mm) == 1
