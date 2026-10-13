from datetime import date, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.domain.kpi import pmo_monthly_kpi_values
from app.models.core import Personnel, Role, ProjectAssignment, MonthlyEmploymentMM, CurrentAssignmentSnapshot, JobLock
from app.services import personnel_excel as excel
from test_capabilities import capability_db as capability_db, headers, grant
from test_project_excel import dataset, seed as seed_projects


@pytest.fixture(autouse=True)
def cache(monkeypatch):
    monkeypatch.setattr(excel, "VALIDATIONS", {})


def seed(factory):
    with factory() as session:
        session.add(Role(id="sales-role", code="SALES_OWNER", name="영업대표", is_active=True))
        for kind in ("pmo", "sales", "technical", "external"):
            values = dict(id=str(uuid4()), personnel_type=kind, name=kind, group_name="PMO본부" if kind == "pmo" else kind, position_name="책임", employment_status="active", is_active=True)
            if kind == "pmo":
                values.update(employee_no="RAW-007", email="pmo@test.local", team_name="운영팀", mm_start_date=date(date.today().year, 1, 1))
            if kind == "sales":
                values.update(role_id="sales-role", role_name="영업대표")
            session.add(Personnel(**values))
        session.commit()
        return {p.personnel_type.value: p.id for p in session.scalars(select(Personnel)).all()}


def export(client, ids=None):
    response = client.post("/api/personnel-excel/export", headers=headers(), json={"ids": ids})
    assert response.status_code == 200, response.text
    return response.content


def edit(raw, changes=None, new=None):
    workbook = load_workbook(BytesIO(raw))
    sheet = workbook.active
    columns = {cell.value: cell.column for cell in sheet[2]}
    for number, values in (changes or {}).items():
        for label, value in values.items():
            sheet.cell(number, columns[label]).value = value
    for values in new or []:
        sheet.append([values.get(label) for label, _ in excel.COLUMNS])
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


def new_row(kind, name):
    values = {"인력 유형": kind, "성명": name, "본부": kind, "직위": "책임", "사용여부": "사용"}
    if kind == "sales":
        values["역할 코드"] = "SALES_OWNER"
    if kind == "pmo":
        values.update({"본부": "PMO본부", "팀": "운영팀", "사번": name, "이메일": f"{name}@test.local", "MM 시작일": f"{date.today().year}-01-01"})
    return values


def validate(client, raw, kind=None, actor="developer"):
    response = client.post("/api/personnel-excel/validate", params={"personnel_type": kind} if kind else None, headers=headers(actor), content=raw)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def apply(client, preview, actions=None, actor="developer", **extra):
    return client.post("/api/personnel-excel/apply", headers=headers(actor), json={"validation_id": preview["validation_id"], "file_sha256": preview["file_sha256"], "missing_actions": actions or {}, **extra})


def test_list_type_scope_roundtrip_and_raw_identifiers(capability_db):
    client, factory = capability_db
    ids = seed(factory)
    before = dataset(factory)
    assert client.get("/api/personnel").json()["meta"]["total"] == 4
    for kind, identity in ids.items():
        rows = client.get("/api/personnel", params={"personnel_type": kind}).json()["data"]
        assert [row["id"] for row in rows] == [identity]
    assert client.get("/api/personnel", params={"personnel_type": "invalid"}).status_code == 422
    assert client.get("/api/personnel/sales-owner-candidates").json()["data"][0]["id"] == ids["sales"]
    raw = export(client)
    sheet = load_workbook(BytesIO(raw)).active
    columns = {cell.value: cell.column for cell in sheet[2]}
    pm = next(row for row in sheet.iter_rows(min_row=3) if row[columns["personnel_id"] - 1].value == ids["pmo"])
    assert pm[columns["사번"] - 1].value == "RAW-007"
    template = client.get("/api/personnel-excel/template", headers=headers())
    assert [c.value for c in load_workbook(BytesIO(template.content)).active[2]] == [c.value for c in sheet[2]]
    preview = validate(client, raw)
    assert preview["summary"] == {"new": 0, "update": 0, "unchanged": 4, "error": 0}
    assert dataset(factory) == before
    for _ in range(2):
        assert apply(client, preview).status_code == 200
        assert dataset(factory) == before
        preview = validate(client, raw)


def test_all_four_new_types_derived_values_uuid_and_backup_recovery(capability_db):
    client, factory = capability_db
    seed(factory)
    before = dataset(factory)
    template = client.get("/api/personnel-excel/template", headers=headers()).content
    raw = edit(template, new=[new_row(kind, "new_" + kind) for kind in ("pmo", "sales", "technical", "external")])
    preview = validate(client, raw)
    assert preview["summary"]["new"] == 4 and preview["summary"]["error"] == 0, preview["issues"]
    response = apply(client, preview)
    assert response.status_code == 200, response.text
    assert len(response.json()["data"]["created_ids"]) == 4
    with factory() as session:
        assert session.execute(text("PRAGMA foreign_key_check")).all() == []
        for person in session.scalars(select(Personnel).where(Personnel.name.like("new_%"))).all():
            mm = session.scalars(select(MonthlyEmploymentMM).where(MonthlyEmploymentMM.personnel_id == person.id)).all()
            assert bool(mm) == (person.personnel_type == "pmo")
            assert (person.yearly_mm is not None) == (person.personnel_type == "pmo")
            assert all(row.month >= date.today().month for row in mm)
    again = validate(client, export(client))
    assert again["summary"]["error"] == 0 and again["summary"]["unchanged"] == 8
    duplicate = validate(client, raw)
    assert duplicate["summary"]["error"] == 4  # blank IDs cannot silently recreate existing people
    backup_id = response.json()["data"]["pre_backup_id"]
    assert client.post("/api/data-backup/restore", headers=headers(), json={"backup_id": backup_id}).status_code == 200
    assert dataset(factory) == before


def test_missing_scope_deactivation_type_change_and_history_preservation(capability_db):
    client, factory = capability_db
    ids = seed(factory)
    project_ids = seed_projects(factory, 1)
    yesterday = date.today() - timedelta(days=1)
    with factory() as session:
        session.add(ProjectAssignment(id="pmo-assignment", personnel_id=ids["pmo"], project_id=project_ids[0], assignment_type="delivery"))
        session.add(MonthlyEmploymentMM(id="history-mm", personnel_id=ids["pmo"], year=2020, month=1, employment_mm=1))
        session.add(CurrentAssignmentSnapshot(id="past-snapshot", personnel_id=ids["pmo"], as_of_date=yesterday, representative_status="delivery", project_id=project_ids[0]))
        session.commit()
    before = dataset(factory)
    raw = export(client, [ids["pmo"]])
    preview = validate(client, edit(raw, {3: {"성명": "renamed", "인력 유형": "technical", "연간 재직 MM": "=999999"}}), "pmo")
    assert preview["summary"]["error"] == 0 and preview["missing"] == []
    assert apply(client, preview, {ids["external"]: "deactivate"}).status_code == 400
    assert apply(client, preview).status_code == 200
    with factory() as session:
        assert session.get(Personnel, ids["pmo"]).personnel_type.value == "technical"
        assert session.get(MonthlyEmploymentMM, "history-mm").employment_mm == 1
        assert pmo_monthly_kpi_values(session, 2020, 1)["avg_headcount_mm"] == 0
        assert session.get(CurrentAssignmentSnapshot, "past-snapshot") is not None
    after = dataset(factory)
    for key in ("project_assignments", "projects", "project_logs", "monthly_assignment_mm", "monthly_kpi_summaries"):
        assert after[key] == before[key]
    # A partial technical file only offers technical rows as missing.
    preview = validate(client, export(client, [ids["technical"]]), "technical")
    assert ids["sales"] not in {r["code"] for r in preview["missing"]}
    assert apply(client, preview, {r["code"]: "deactivate" for r in preview["missing"]}).status_code == 200
    with factory() as session:
        assert session.get(Personnel, ids["pmo"]).is_active is False
        assert session.get(Personnel, ids["sales"]).is_active is True


@pytest.mark.parametrize("changes", [
    {"인력 유형": None}, {"인력 유형": "bad"}, {"personnel_id": "unknown"}, {"성명": None},
    {"성명": "=1+1"}, {"이메일": "bad"}, {"사번": None}, {"팀": None}, {"직위": None},
    {"MM 시작일": None}, {"MM 종료일": "1999-01-01"}, {"MM 시작일": "bad"},
    {"사용여부": "bad"}, {"재직상태": "bad"}, {"역할 코드": "bad"}, {"역할 코드": "SALES_OWNER"},
])
def test_invalid_input_is_preview_error_without_writes(capability_db, changes):
    client, factory = capability_db
    ids = seed(factory)
    before = dataset(factory)
    preview = validate(client, edit(export(client, [ids["pmo"]]), {3: changes}))
    assert preview["summary"]["error"] == 1
    assert apply(client, preview).status_code == 409
    assert dataset(factory) == before


def test_tokens_permissions_db_changes_duplicates_and_failure_rollback(capability_db, monkeypatch):
    client, factory = capability_db
    ids = seed(factory)
    for route, method, body in [("template", "GET", None), ("export", "POST", {}), ("validate", "POST", None), ("apply", "POST", {"validation_id": "none", "file_sha256": "none"})]:
        assert client.request(method, f"/api/personnel-excel/{route}", headers=headers("admin"), json=body).status_code == 403
        assert client.request(method, f"/api/personnel-excel/{route}", json=body).status_code == 401
    grant(client, "member", ["personnel_excel_manage"])
    raw = edit(export(client, [ids["technical"]]), {3: {"직위": "이사"}}, new=[new_row("external", "brand_new")])
    preview = validate(client, raw)
    assert preview["summary"]["error"] == 0
    assert apply(client, preview, actor="member").status_code == 403
    assert apply(client, preview, file_sha256="wrong").status_code == 409
    before = dataset(factory)
    original = excel.refresh_personnel_projections
    calls = []

    def fail_after_one(session, person):
        calls.append(person.id)
        original(session, person)
        if len(calls) == 1:
            raise RuntimeError("simulated derived failure")

    monkeypatch.setattr(excel, "refresh_personnel_projections", fail_after_one)
    with pytest.raises(RuntimeError, match="simulated derived failure"):
        apply(client, preview)
    assert dataset(factory) == before
    with factory() as session:
        assert session.scalars(select(JobLock)).all() == []
    monkeypatch.setattr(excel, "refresh_personnel_projections", original)
    with factory() as session:
        session.get(Role, "sales-role").name = "changed role"
        session.commit()
    assert apply(client, preview).status_code == 409
    preview = validate(client, raw, actor="member")
    grant(client, "member", [])
    assert apply(client, preview, actor="member").status_code == 403
    preview = validate(client, raw)
    excel.VALIDATIONS[preview["validation_id"]]["created_at"] -= timedelta(minutes=31)
    assert apply(client, preview).status_code == 409


def test_uuid_and_employee_collisions_sales_role_and_old_template_rejection(capability_db):
    client, factory = capability_db
    ids = seed(factory)
    row = new_row("external", "collision") | {"사번": "RAW-007"}
    template = client.get("/api/personnel-excel/template", headers=headers()).content
    assert validate(client, edit(template, new=[row]))["summary"]["error"] == 1
    assert validate(client, edit(template, new=[new_row("sales", "no_role") | {"역할 코드": None}]))["summary"]["error"] == 1
    assert validate(client, edit(template, new=[new_row("sales", "pmo")]))["summary"]["error"] == 1
    raw = export(client, [ids["pmo"]])
    sheet = load_workbook(BytesIO(raw)).active
    duplicate = {cell.value: sheet.cell(3, cell.column).value for cell in sheet[2]}
    assert validate(client, edit(raw, new=[duplicate]))["summary"]["error"] >= 1
    workbook = load_workbook(BytesIO(raw))
    workbook.active.cell(2, 1).value = "old ID"
    result = BytesIO()
    workbook.save(result)
    assert client.post("/api/personnel-excel/validate", headers=headers(), content=result.getvalue()).status_code == 400


def test_basic_edit_preserves_manual_current_mm_and_projection_edit_preserves_notes(capability_db):
    client, factory = capability_db
    ids = seed(factory)
    today = date.today()
    with factory() as session:
        session.add(MonthlyEmploymentMM(id="manual-mm", personnel_id=ids["pmo"], year=today.year, month=today.month, employment_mm=0.125, note="수동 보정"))
        session.add(CurrentAssignmentSnapshot(id="current-note", personnel_id=ids["pmo"], as_of_date=today, representative_status="unassigned", weekly_note="운영 메모", monthly_mm={"manual": 0.125}))
        session.commit()
    preview = validate(client, edit(export(client, [ids["pmo"]]), {3: {"직위": "이사", "연간 재직 MM": 99999}}))
    assert apply(client, preview).status_code == 200
    with factory() as session:
        assert float(session.get(MonthlyEmploymentMM, "manual-mm").employment_mm) == 0.125
        assert session.get(CurrentAssignmentSnapshot, "current-note").weekly_note == "운영 메모"
    preview = validate(client, edit(export(client, [ids["pmo"]]), {3: {"MM 종료일": today.isoformat()}}))
    assert apply(client, preview).status_code == 200
    with factory() as session:
        row = session.get(CurrentAssignmentSnapshot, "current-note")
        assert row and row.weekly_note == "운영 메모" and row.monthly_mm == {"manual": 0.125}
        assert session.get(MonthlyEmploymentMM, "manual-mm").note == "수동 보정"


def test_typed_crud_required_fields_and_legacy_missing_type_compatibility(capability_db):
    client, factory = capability_db
    seed(factory)
    assert client.post("/api/personnel", json={"personnel_type": "pmo", "name": "incomplete", "group_name": "PMO본부"}).status_code == 400
    old = client.post("/api/personnel", json={"name": "legacy", "group_name": "PMO본부"})
    assert old.status_code == 201
    assert old.json()["data"]["personnel_type"] == "pmo"
    assert client.patch(f"/api/personnel/{old.json()['data']['id']}", json={"note": "omit type"}).json()["data"]["personnel_type"] == "pmo"


def test_backup_failure_or_duplicate_open_ledger_rolls_back(capability_db, monkeypatch):
    from app.services import data_backup as backup
    client, factory = capability_db
    ids = seed(factory)
    preview = validate(client, edit(export(client, [ids["pmo"]]), {3: {"MM 종료일": date.today().isoformat()}}))
    before = dataset(factory)
    original = backup.create_backup
    monkeypatch.setattr(backup, "create_backup", lambda *args, **kwargs: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError, match="disk full"):
        apply(client, preview)
    assert dataset(factory) == before
    monkeypatch.setattr(backup, "create_backup", original)
    with factory() as session:
        for _ in range(2):
            session.add(MonthlyEmploymentMM(personnel_id=ids["pmo"], year=date.today().year, month=date.today().month, employment_mm=1))
        session.commit()
    preview = validate(client, edit(export(client, [ids["pmo"]]), {3: {"MM 종료일": date.today().isoformat()}}))
    before = dataset(factory)
    assert apply(client, preview).status_code == 409
    assert dataset(factory) == before
    with factory() as session:
        assert session.scalars(select(JobLock)).all() == []
