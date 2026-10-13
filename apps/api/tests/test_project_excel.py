from datetime import date, datetime, timedelta
from io import BytesIO
import json
from urllib.parse import quote

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.models.core import (
    CurrentAssignmentSnapshot, JobLock, MonthlyAssignmentMM, MonthlyKpiSummary,
    Personnel, Project, ProjectAssignment, ProjectCode, ProjectLog, User, WeeklyLoadSnapshot,
)
from app.schemas.projects import ProjectMasterCreate
from app.services import data_backup as backup, project_excel as excel
from app.services.project_master import create_project_master
from test_capabilities import capability_db as capability_db, grant, headers


@pytest.fixture(autouse=True)
def clear_validation_cache(monkeypatch):
    monkeypatch.setattr(excel, "VALIDATIONS", {})


def project_values(code="P2026001", name="기존 사업"):
    return dict(code=code, name=name, client_name="고객사", project_type="main", status="running",
                certainty="확정", total_amount=123456789.25, company_amount=23456789.25,
                sales_department="영업본부", sales_owner="영업대표", delivery_pm_name="수행PM",
                start_date=date(2026, 1, 1), end_date=date(2026, 12, 31))


def seed(factory, count=3):
    with factory() as session:
        user = session.get(User, "developer")
        projects = []
        for index in range(1, count + 1):
            values = project_values(f"P2026{index:03}", f"기존 사업 {index}")
            _, project = create_project_master(session, user, ProjectMasterCreate(
                project_code={key: values[key] for key in ("code", "name", "project_type", "status", "certainty")}, project=values,
            ))
            projects.append(project)
        session.add(Personnel(id="technical", name="기술인력", position_name="책임", personnel_type="technical", employment_status="active", is_active=True))
        session.flush()
        session.add(ProjectAssignment(id="assignment", project_id=projects[0].id, personnel_id="technical", assignment_type="delivery", monthly_mm={"2026-01": 1}, start_date=date(2026, 1, 1), end_date=date(2026, 12, 31)))
        session.flush()
        session.add(MonthlyAssignmentMM(id="mm", assignment_id="assignment", project_id=projects[0].id, personnel_id="technical", year=2026, month=1, assignment_mm=1, assignment_type="delivery"))
        session.add(MonthlyKpiSummary(id="kpi", year=2026, month=1, organization_name="PMO본부", avg_headcount_mm=1))
        for snapshot_id, day in [("historical", date(2026, 1, 31)), ("current", date.today())]:
            session.add(CurrentAssignmentSnapshot(id=snapshot_id, as_of_date=day, personnel_id="technical", representative_status="running", project_id=projects[0].id, project_name=projects[0].name, project_code=projects[0].code, assignment_id="assignment"))
            session.add(WeeklyLoadSnapshot(id=f"week-{snapshot_id}", as_of_date=day, personnel_id="technical", week_offset=0, week_label="W0", representative_status="running", project_id=projects[0].id, project_name=projects[0].name))
        session.commit()
        return [p.id for p in projects]


def download(client, codes=None):
    response = client.post("/api/project-excel/export", headers=headers(), json={"codes": codes})
    assert response.status_code == 200, response.text
    return response.content


def edit_workbook(raw, edits=None, append=None):
    workbook = load_workbook(BytesIO(raw))
    sheet = workbook[excel.SHEET]
    columns = {cell.value: cell.column for cell in sheet[2]}
    for row_number, values in (edits or {}).items():
        for label, value in values.items():
            sheet.cell(row_number, columns[label]).value = value
    for values in append or []:
        sheet.append([values.get(label) for label, _ in excel.COLUMNS])
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


def new_row(code="P2026099"):
    values = project_values(code, "신규 사업")
    return {label: values.get(field) for label, field in excel.COLUMNS} | {"사용여부": "사용"}


def validate(client, raw, actor="developer"):
    response = client.post("/api/project-excel/validate", headers=headers(actor, **{"x-upload-filename": quote("프로젝트.xlsx")}), content=raw)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def apply(client, preview, actions=None, actor="developer", **overrides):
    payload = {"validation_id": preview["validation_id"], "file_sha256": preview["file_sha256"], "missing_actions": actions or {}}
    return client.post("/api/project-excel/apply", headers=headers(actor), json=payload | overrides)


def dataset(factory):
    with factory() as session:
        return {key: sorted(backup._serialize_table_rows(session, model), key=lambda row: row["id"]) for key, model in backup.TABLE_MODELS.items()}


def test_template_export_and_unchanged_round_trip_are_read_only(capability_db):
    client, factory = capability_db
    ids = seed(factory)
    with factory() as session:
        project = session.get(Project, ids[0])
        project.submission_note = "=literal business text"
        project.submission_at = datetime(2026, 3, 4, 15, 16, 17, 123456)
        session.commit()
    before = dataset(factory)
    template = client.get("/api/project-excel/template", headers=headers())
    assert template.status_code == 200
    empty = load_workbook(BytesIO(template.content))
    raw = download(client)
    workbook = load_workbook(BytesIO(raw))
    sheet = workbook[excel.SHEET]
    assert [c.value for c in empty.active[2]] == [c.value for c in sheet[2]]
    assert empty.active.max_row == 2 and sheet.max_row == 5 and sheet.freeze_panes == "C3"
    assert len(sheet.merged_cells.ranges) == 5
    columns = {cell.value: cell.column for cell in sheet[2]}
    assert "사업금액" not in columns
    assert sheet.cell(5, columns["총 사업금액"]).value == 123456789.25
    assert sheet.cell(5, columns["수행PM"]).value == "수행PM"
    assert sheet.cell(5, columns["제안/수행팀"]).value == "기술인력 책임"
    assert sheet.cell(5, columns["제출 유의사항"]).data_type == "s"
    preview = validate(client, raw)
    assert preview["summary"] == {"new": 0, "update": 0, "unchanged": 3, "error": 0}
    assert dataset(factory) == before and preview["missing"] == []
    for _ in range(2):
        result = apply(client, preview)
        assert result.status_code == 200, result.text
        assert result.json()["data"]["counts"]["unchanged"] == 3
        assert dataset(factory) == before
        preview = validate(client, raw)
    with factory() as session:
        assert session.get(Project, ids[0]).submission_at.microsecond == 123456
    assert client.post("/api/project-excel/export", headers=headers(), json={"codes": ["gone"]}).status_code == 409


def test_upsert_preserves_ids_references_ledgers_and_historical_snapshots(capability_db):
    client, factory = capability_db
    ids = seed(factory)
    before = dataset(factory)
    raw = edit_workbook(download(client, ["P2026001"]), {3: {"사업명": "수정 사업", "당사 사업금액": 34567890.25, "제안/수행팀": "위조팀", "최근 활동일": "위조일"}}, [new_row()])
    preview = validate(client, raw)
    assert preview["summary"] == {"new": 1, "update": 1, "unchanged": 0, "error": 0}
    assert [r["code"] for r in preview["missing"]] == ["P2026002", "P2026003"]
    assert dataset(factory) == before
    result = apply(client, preview)
    assert result.status_code == 200, result.text
    assert result.json()["data"]["counts"] == {"new": 1, "update": 1, "unchanged": 0, "deactivated": 0}
    with factory() as session:
        project = session.get(Project, ids[0])
        assert project.name == project.project_code.name == "수정 사업"
        assert project.project_code.id == before["projects"][next(i for i, row in enumerate(before["projects"]) if row["id"] == ids[0])]["project_code_id"]
        assert float(project.company_amount) == 34567890.25 and project.amount_text == "1.2억/3457만"
        assert project.recent_activity_at == session.scalar(select(ProjectLog.logged_at).where(ProjectLog.project_id == project.id).order_by(ProjectLog.logged_at.desc()))
        assert session.get(CurrentAssignmentSnapshot, "current").project_name == "수정 사업"
        assert session.get(WeeklyLoadSnapshot, "week-current").project_name == "수정 사업"
        assert session.get(CurrentAssignmentSnapshot, "historical").project_name == "기존 사업 1"
        assert session.get(WeeklyLoadSnapshot, "week-historical").project_name == "기존 사업 1"
        assert all(p.is_active for p in session.scalars(select(ProjectCode)).all())
        assert session.execute(text("PRAGMA foreign_key_check")).all() == []
    after = dataset(factory)
    for table in ["personnel", "project_assignments", "monthly_assignment_mm", "monthly_kpi_summaries"]:
        assert after[table] == before[table]
    old_logs = {row["id"]: row for row in before["project_logs"]}
    assert all(row == old_logs[row["id"]] for row in after["project_logs"] if row["id"] in old_logs)
    backup_id = result.json()["data"]["pre_backup_id"]
    saved = json.loads((backup.BACKUP_DIR / f"{backup_id}.json").read_text(encoding="utf-8"))
    assert saved["kind"] == "project_upload_before_backup"
    assert sorted(saved["tables"]["projects"], key=lambda row: row["id"]) == before["projects"]
    assert apply(client, preview).status_code == 409
    # Recovery uses the existing integrated backup/restore API on this temporary DB.
    restored = client.post("/api/data-backup/restore", headers=headers(), json={"backup_id": backup_id, "memo": "test recovery"})
    assert restored.status_code == 200, restored.text
    assert dataset(factory) == before


def test_missing_deactivation_is_explicit_and_does_not_delete_assignments(capability_db):
    client, factory = capability_db
    ids = seed(factory)
    preview = validate(client, download(client, ["P2026002"]))
    assert apply(client, preview, {"P2026002": "deactivate"}).status_code == 400
    response = apply(client, preview, {"P2026001": "deactivate"})
    assert response.status_code == 200, response.text
    with factory() as session:
        assert session.get(Project, ids[0]).project_code.is_active is False
        assert session.get(ProjectAssignment, "assignment").project_id == ids[0]
        assert session.get(MonthlyAssignmentMM, "mm") is not None
        assert session.get(Project, ids[2]).project_code.is_active is True
    preview = validate(client, download(client, ["P2026002"]))
    assert apply(client, preview, {row["code"]: "deactivate" for row in preview["missing"]}).json()["data"]["counts"]["deactivated"] == 1


@pytest.mark.parametrize("changes", [
    {"코드": "invalid"}, {"총 사업금액": -1}, {"총 사업금액": "NaN"},
    {"총 사업금액": "100000000000000"}, {"당사 사업금액": "1.001"}, {"당사 사업금액": 999999999},
    {"사업 시작일": "bad"}, {"사업 종료일": "2025-01-01"}, {"수행PM": None},
    {"사업유형": "invalid"}, {"상태": "loss"}, {"사용여부": "unknown"},
    {"사업명": "=1+1"}, {"고객사": None}, {"제안PM": "제안 담당"},
    {"제안서 제출일": "2026-01-01T00:00:00+09:00"},
])
def test_validation_errors_block_all_writes(capability_db, changes):
    client, factory = capability_db
    seed(factory, 1)
    before = dataset(factory)
    preview = validate(client, edit_workbook(download(client), {3: changes}))
    assert preview["summary"]["error"] == 1 and preview["issues"]
    assert apply(client, preview).status_code == 409
    assert dataset(factory) == before and not backup.MANIFEST_PATH.exists()


def test_duplicate_codes_wrong_sheets_headers_empty_and_legacy_format(capability_db):
    client, factory = capability_db
    seed(factory, 1)
    raw = download(client)
    duplicate = validate(client, edit_workbook(raw, append=[new_row("P2026001")]))
    assert duplicate["summary"]["error"] == 1
    assert client.post("/api/project-excel/validate", headers=headers(), content=b"invalid").status_code == 400
    template = client.get("/api/project-excel/template", headers=headers()).content
    assert client.post("/api/project-excel/validate", headers=headers(), content=template).status_code == 400
    for kind in ["sheet", "duplicate_header", "missing_header"]:
        workbook = load_workbook(BytesIO(raw))
        if kind == "sheet":
            workbook.create_sheet("인력관리")
        else:
            workbook.active.cell(2, 27, "코드") if kind == "duplicate_header" else setattr(workbook.active.cell(2, 2), "value", "다른 항목")
        stream = BytesIO()
        workbook.save(stream)
        assert client.post("/api/project-excel/validate", headers=headers(), content=stream.getvalue()).status_code == 400
    workbook = load_workbook(BytesIO(raw))
    sheet = workbook.active
    # Remove group merges before converting an old 25-column combined-amount file.
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    sheet.cell(2, 6, "사업금액")
    sheet.cell(3, 6, "123456789.25/23456789.25")
    sheet.delete_cols(7)
    stream = BytesIO()
    workbook.save(stream)
    preview = validate(client, stream.getvalue())
    assert preview["summary"]["unchanged"] == 1 and len(preview["warnings"]) == 2
    sheet.cell(3, 6, "nan/1")
    stream = BytesIO()
    workbook.save(stream)
    assert validate(client, stream.getvalue())["summary"]["error"] == 1


def test_permission_user_hash_expiry_and_db_change_binding(capability_db):
    client, factory = capability_db
    ids = seed(factory, 1)
    for method, route, body in [("GET", "template", None), ("POST", "export", {}), ("POST", "validate", None), ("POST", "apply", {"validation_id": "none", "file_sha256": "none"})]:
        assert client.request(method, f"/api/project-excel/{route}", json=body).status_code == 401
        assert client.request(method, f"/api/project-excel/{route}", headers=headers("admin", **{"x-user-capabilities": "project_excel_manage"}), json=body).status_code == 403
    assert grant(client, "member", ["project_excel_manage"]).status_code == 200
    raw = edit_workbook(download(client), {3: {"사업명": "수정"}})
    preview = validate(client, raw)
    assert apply(client, preview, actor="member").status_code == 403
    assert apply(client, preview, file_sha256="other").status_code == 409
    assert apply(client, preview, {"P2026001": "delete"}).status_code == 422
    with factory() as session:
        session.get(Project, ids[0]).client_name = "다른 사용자 수정"
        session.commit()
    assert apply(client, preview).status_code == 409
    preview = validate(client, download(client))
    excel.VALIDATIONS[preview["validation_id"]]["created_at"] -= timedelta(minutes=31)
    assert apply(client, preview).status_code == 409
    member_preview = validate(client, download(client), "member")
    assert grant(client, "member", []).status_code == 200
    assert apply(client, member_preview, actor="member").status_code == 403
    assert not backup.MANIFEST_PATH.exists()


def test_transaction_rollback_lock_release_and_retry(capability_db, monkeypatch):
    client, factory = capability_db
    seed(factory)
    before = dataset(factory)
    raw = edit_workbook(download(client, ["P2026001"]), {3: {"사업명": "롤백할 수정"}}, [new_row()])
    preview = validate(client, raw)
    original = excel.create_project_master

    def fail_after_update(*args, **kwargs):
        raise RuntimeError("simulated import failure after first project flush")

    monkeypatch.setattr(excel, "create_project_master", fail_after_update)
    with pytest.raises(RuntimeError, match="simulated import failure"):
        apply(client, preview, {"P2026002": "deactivate"})
    assert dataset(factory) == before
    with factory() as session:
        assert session.scalars(select(JobLock)).all() == []
        assert session.execute(text("PRAGMA foreign_key_check")).all() == []
    monkeypatch.setattr(excel, "create_project_master", original)
    assert apply(client, preview).status_code == 200


def test_bulk_lock_and_backup_failure_block_writes(capability_db, monkeypatch):
    client, factory = capability_db
    seed(factory, 1)
    preview = validate(client, edit_workbook(download(client), {3: {"사업명": "수정"}}))
    before = dataset(factory)
    with factory() as session:
        backup._acquire_lock(session, locked_by="restore worker")
    assert apply(client, preview).status_code == 409
    with factory() as session:
        backup._release_lock(session)

    def fail_backup(*args, **kwargs):
        raise OSError("simulated full backup disk")

    monkeypatch.setattr(backup, "create_backup", fail_backup)
    with pytest.raises(OSError, match="simulated full backup disk"):
        apply(client, preview)
    assert dataset(factory) == before
    with factory() as session:
        assert session.scalars(select(JobLock)).all() == []


def test_changed_status_uses_existing_transition_log_and_empty_active_preserves(capability_db):
    client, factory = capability_db
    ids = seed(factory, 1)
    with factory() as session:
        session.get(Project, ids[0]).project_code.is_active = False
        session.commit()
    preview = validate(client, edit_workbook(download(client), {3: {"상태": "done", "사용여부": None}}))
    assert preview["summary"]["update"] == 1
    assert apply(client, preview).status_code == 200
    with factory() as session:
        assert session.get(Project, ids[0]).status.value == "done"
        assert session.get(Project, ids[0]).project_code.is_active is False
        log = session.scalar(select(ProjectLog).where(ProjectLog.next_status == "done"))
        assert log.previous_status.value == "running"


def test_link_inconsistency_is_not_silently_repaired(capability_db):
    client, factory = capability_db
    ids = seed(factory, 1)
    raw = download(client)
    with factory() as session:
        session.get(Project, ids[0]).project_code.name = "코드와 프로젝트 불일치"
        session.commit()
    assert client.post("/api/project-excel/export", headers=headers(), json={}).status_code == 409
    assert client.post("/api/project-excel/validate", headers=headers(), content=raw).status_code == 409


def test_legacy_rounded_display_never_replaces_precise_stored_amounts(capability_db):
    client, factory = capability_db
    ids = seed(factory, 1)
    with factory() as session:
        session.get(Project, ids[0]).amount_text = "1.2억/2346만"
        session.commit()
    workbook = load_workbook(BytesIO(download(client)))
    sheet = workbook.active
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    sheet.cell(2, 6, "사업금액")
    sheet.cell(3, 6, "1.2억/2346만")
    sheet.delete_cols(7)
    stream = BytesIO()
    workbook.save(stream)
    preview = validate(client, stream.getvalue())
    assert preview["summary"]["unchanged"] == 1 and preview["summary"]["error"] == 0


def test_concurrent_change_during_validation_requires_revalidation(capability_db, monkeypatch):
    client, factory = capability_db
    seed(factory, 1)
    raw = download(client)
    original = excel._fingerprint
    reads = []

    def changed_fingerprint(session):
        reads.append(True)
        return original(session) if len(reads) == 1 else "concurrently changed"

    monkeypatch.setattr(excel, "_fingerprint", changed_fingerprint)
    assert client.post("/api/project-excel/validate", headers=headers(), content=raw).status_code == 409
    assert excel.VALIDATIONS == {}


def test_same_second_metadata_change_and_oversize_file_are_rejected(capability_db, monkeypatch):
    client, factory = capability_db
    ids = seed(factory, 1)
    with factory() as session:
        session.get(Project, ids[0]).updated_at = datetime(2026, 1, 1, 0, 0, 0, 123456)
        session.commit()
    preview = validate(client, download(client))
    with factory() as session:
        session.get(Project, ids[0]).updated_at = datetime(2026, 1, 1, 0, 0, 0, 234567)
        session.commit()
    assert apply(client, preview).status_code == 409
    monkeypatch.setattr(excel, "MAX_FILE_BYTES", 5)
    assert client.post("/api/project-excel/validate", headers=headers(), content=b"123456").status_code == 400
